"""Sotto-router Statistiche Produzione — export xlsx/csv/pdf.

Montato sotto /produzione dall'aggregatore produzione.py. Self-contained (non riusa
app/routers/export.py: quelle funzioni sono private e cucite sui modelli Revenue).
"""
import csv
import io
from datetime import date
from typing import Optional

import openpyxl
from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from app.auth import richiedi_utente_attivo
from app.database import get_db
from app.models.produzione import ProdCategoria
from app.routers.produzione_report import _categorie_attive, _per_dimensione, report_giornaliero, report_mensile

router = APIRouter(dependencies=[Depends(richiedi_utente_attivo)])

_BLU = "1e3a5f"
_GRIGIO = "f1f5f9"


def _xlsx(intestazioni: list[str], righe: list[list]) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(intestazioni)
    for c in range(1, len(intestazioni) + 1):
        cella = ws.cell(row=1, column=c)
        cella.font = Font(bold=True, color="FFFFFF")
        cella.fill = PatternFill("solid", fgColor=_BLU)
        cella.alignment = Alignment(horizontal="center")
    for i, riga in enumerate(righe, start=2):
        ws.append(riga)
        if i % 2 == 0:
            for c in range(1, len(intestazioni) + 1):
                ws.cell(row=i, column=c).fill = PatternFill("solid", fgColor=_GRIGIO)
    for c in range(1, len(intestazioni) + 1):
        letter = get_column_letter(c)
        larghezza = max(12, min(30, max(len(str(r[c - 1])) for r in [intestazioni] + righe) + 2))
        ws.column_dimensions[letter].width = larghezza
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _csv(intestazioni: list[str], righe: list[list]) -> io.BytesIO:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=';')
    w.writerow(intestazioni)
    w.writerows(righe)
    return io.BytesIO(buf.getvalue().encode('utf-8-sig'))


def _pdf(intestazioni: list[str], righe: list[list], titolo: str) -> io.BytesIO:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), topMargin=20, bottomMargin=20)
    stili = getSampleStyleSheet()
    elementi = [Paragraph(titolo, stili['Heading2']), Spacer(1, 10)]
    dati = [intestazioni] + [[str(x) for x in r] for r in righe]
    tabella = Table(dati, repeatRows=1)
    tabella.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(f'#{_BLU}')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor(f'#{_GRIGIO}')]),
    ]))
    elementi.append(tabella)
    doc.build(elementi)
    buf.seek(0)
    return buf


def _risposta(formato: str, nome: str, titolo: str, intestazioni: list[str], righe: list[list]) -> StreamingResponse:
    if formato == 'csv':
        buf, media = _csv(intestazioni, righe), 'text/csv'
    elif formato == 'pdf':
        buf, media = _pdf(intestazioni, righe, titolo), 'application/pdf'
    else:
        buf, media = _xlsx(intestazioni, righe), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return StreamingResponse(buf, media_type=media, headers={
        'Content-Disposition': f'attachment; filename="{nome}.{formato}"',
    })


@router.get("/export/mensile")
def export_mensile(
    anno: int = Query(...),
    struttura_code: Optional[str] = Query(None),
    formato: str = Query('xlsx'),
    is_test: bool = Query(False),
    db: Session = Depends(get_db),
):
    dati = report_mensile(anno=anno, mese=None, struttura_code=struttura_code, canale=None,
                          trattamento=None, is_test=is_test, db=db)
    from app.utils.locale_it import MESI_IT
    nomi_mesi = {v: k.capitalize() for k, v in MESI_IT.items()}

    intestazioni = ['Mese'] + list(dati['mesi'][0]['per_struttura'].keys()) + ['Totale']
    righe = []
    for mo in dati['mesi']:
        riga = [nomi_mesi[mo['mese']]]
        for sc in dati['mesi'][0]['per_struttura'].keys():
            riga.append(mo['per_struttura'][sc]['totale']['lordo'])
        riga.append(mo['totale']['lordo'])
        righe.append(riga)
    righe.append(['ANNO'] + [
        sum(mo['per_struttura'][sc]['totale']['lordo'] for mo in dati['mesi'])
        for sc in dati['mesi'][0]['per_struttura'].keys()
    ] + [dati['totale_anno']['lordo']])

    return _risposta(formato, f'produzione_mensile_{anno}', f'Produzione — Riepilogo {anno}', intestazioni, righe)


@router.get("/export/giornaliero")
def export_giornaliero(
    data_da: date = Query(...),
    data_a: date = Query(...),
    struttura_code: Optional[str] = Query(None),
    categoria_code: Optional[str] = Query(None),
    canale: Optional[str] = Query(None),
    trattamento: Optional[str] = Query(None),
    formato: str = Query('xlsx'),
    is_test: bool = Query(False),
    db: Session = Depends(get_db),
):
    dati = report_giornaliero(
        data_da=data_da, data_a=data_a, struttura_code=struttura_code, categoria_code=categoria_code,
        canale=canale, trattamento=trattamento, is_test=is_test, db=db,
    )
    categorie = _categorie_attive(db)
    intestazioni = ['Data', 'Struttura'] + [c.name for c in categorie] + ['Totale']
    righe = [
        [g['data'], g['struttura_code']] +
        [g['per_categoria'][c.code]['lordo'] for c in categorie] +
        [g['totale']['lordo']]
        for g in dati
    ]
    return _risposta(formato, 'produzione_giornaliero', 'Produzione — Report giornaliero', intestazioni, righe)


@router.get("/export/{dimensione}")
def export_per_dimensione(
    dimensione: str,
    data_da: date = Query(...),
    data_a: date = Query(...),
    struttura_code: Optional[str] = Query(None),
    formato: str = Query('xlsx'),
    is_test: bool = Query(False),
    db: Session = Depends(get_db),
):
    """Export tabellare per canale/trattamento/tipo-ospite (stessa aggregazione dei report).

    ⚠️ Definito DOPO /export/mensile: FastAPI instrada in ordine di registrazione, un path
    parametrizzato registrato prima catturerebbe anche le richieste a /export/mensile.
    """
    campo_map = {'canali': 'canale', 'trattamenti': 'trattamento', 'tipo-ospite': 'tipo_ospite'}
    campo = campo_map.get(dimensione)
    if not campo:
        return {'errore': 'dimensione non valida'}

    dati = _per_dimensione(db, campo, data_da, data_a, struttura_code, is_test)
    intestazioni = [campo.replace('_', ' ').title(), 'Lordo €', 'N. Prenotazioni', 'N. Righe', '% sul Totale']
    righe = [[d[campo], d['lordo'], d['n_prenotazioni'], d['n_righe'], d['pct_totale']] for d in dati]
    return _risposta(formato, f'produzione_{dimensione}', f'Produzione — {dimensione}', intestazioni, righe)
