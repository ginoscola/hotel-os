"""Conto Economico USALI contabile: import dell'export del gestionale + report.

Import in due passi:
- POST /anteprima: parsing + quadratura, file salvati in un'area temporanea con un token,
  confronto netto vecchio→nuovo per i mesi già presenti. Nessuna scrittura sul DB.
- POST /conferma: in un'unica transazione, per ogni (struttura, anno, mese) del file cancella le
  righe esistenti e inserisce le nuove (l'ultimo import vince); file originali conservati su
  disco e import registrato nel log.
"""

import hashlib
import io
import json
import os
import re
import shutil
import time
import uuid
from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Optional

import openpyxl
from fastapi import APIRouter, Body, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import and_, func, or_
from sqlalchemy.orm import Session

from app.auth import richiedi_admin, richiedi_utente_attivo
from app.database import get_db
from app.models.revenue import Hotel, User
from app.models.usali_contabile import (
    UsaliCategoriaKpi, UsaliImport, UsaliRegolaSegmento, UsaliRiga, UsaliSegmento, UsaliStrutturaGestionale,
)
from app.services.usali_contabile_parser import ErroreImportUsali, parse_export
from app.services.usali_contabile_kpi import (
    CATEGORIE_KPI, calcola_kpi, carica_mappa_categorie, categoria_di, produzione_mensile,
)
from app.services.usali_contabile_report import costruisci_report
from app.services.usali_segmenti import assegna_segmenti
from app.utils.locale_it import MESI_IT

router = APIRouter(prefix="/usali/contabile", tags=["usali-contabile"])

STRUTTURE_EXTRA = {"BON": "Buona Onda"}   # strutture non in `hotels` che possono avere un CE
TTL_ANTEPRIMA_SEC = 24 * 3600


def _storage_dir() -> Path:
    """Cartella file originali. Letta a ogni chiamata (i test la ridirigono via env)."""
    default = Path(__file__).resolve().parents[3] / "uploads" / "usali_contabile"
    return Path(os.getenv("USALI_CONTABILE_DIR") or default)


def _nome_sicuro(nome: str) -> str:
    nome = Path(nome or "file").name
    return re.sub(r"[^A-Za-z0-9._ -]", "_", nome)[:150] or "file"


def _strutture_valide(db: Session) -> dict:
    out = {h.code: h.name for h in db.query(Hotel).order_by(Hotel.code).all()}
    out.update(STRUTTURE_EXTRA)
    return out


def _periodo_txt(anno: int, mese: int) -> str:
    return f"{anno:04d}-{mese:02d}"


def _struttura_da_gestionale(db: Session, id_gestionale: Optional[int]) -> Optional[str]:
    if id_gestionale is None:
        return None
    m = db.query(UsaliStrutturaGestionale).filter_by(id_gestionale=id_gestionale).first()
    return m.struttura_code if m else None


def _risolvi_struttura(db: Session, rilevata: Optional[str], richiesta: Optional[str]) -> str:
    validi = _strutture_valide(db)
    if richiesta and richiesta not in validi:
        raise HTTPException(400, f"Struttura '{richiesta}' sconosciuta")
    if rilevata and richiesta and rilevata != richiesta:
        raise HTTPException(
            400, f"Il file è della struttura {rilevata} (id gestionale nel CSV), non di {richiesta}"
        )
    struttura = rilevata or richiesta
    if not struttura:
        raise HTTPException(
            400, "Struttura non rilevata dal file (CSV assente o id gestionale non mappato): sceglierla"
        )
    return struttura


def _pulisci_anteprime_scadute():
    base = _storage_dir() / "_anteprima"
    if not base.exists():
        return
    limite = time.time() - TTL_ANTEPRIMA_SEC
    for d in base.iterdir():
        if d.is_dir() and d.stat().st_mtime < limite:
            shutil.rmtree(d, ignore_errors=True)


def _leggi_anteprima(token: str) -> tuple:
    if not re.fullmatch(r"[0-9a-f]{32}", token or ""):
        raise HTTPException(400, "Token anteprima non valido")
    d = _storage_dir() / "_anteprima" / token
    meta_path = d / "meta.json"
    if not meta_path.exists():
        raise HTTPException(410, "Anteprima scaduta o già confermata: ricaricare i file")
    meta = json.loads(meta_path.read_text())
    xlsx = (d / "file.xlsx").read_bytes()
    csv = (d / "file.csv").read_bytes() if (d / "file.csv").exists() else None
    return d, meta, xlsx, csv


def _import_duplicato(db: Session, struttura: str, sha: str) -> Optional[UsaliImport]:
    return db.query(UsaliImport).filter_by(struttura_code=struttura, sha256=sha).first()


def _confronto_mesi(db: Session, struttura: str, risultato) -> list:
    """Per ogni mese del file già presente in DB: netto e ricavi vecchio → nuovo."""
    periodi = risultato.periodi
    filtro = or_(*[and_(UsaliRiga.anno == a, UsaliRiga.mese == m) for a, m in periodi])
    esistenti = (
        db.query(UsaliRiga.anno, UsaliRiga.mese, UsaliRiga.sezione, func.sum(UsaliRiga.importo))
        .filter(UsaliRiga.struttura_code == struttura, filtro)
        .group_by(UsaliRiga.anno, UsaliRiga.mese, UsaliRiga.sezione)
        .all()
    )
    vecchi = defaultdict(lambda: {"netto": Decimal(0), "ricavi": Decimal(0)})
    for a, m, sez, somma in esistenti:
        vecchi[(a, m)]["netto"] += somma
        if sez == "Ricavi Operativi":
            vecchi[(a, m)]["ricavi"] += somma

    out = []
    for p in periodi:
        nuovo_netto = sum((r.importi[p] for r in risultato.righe), Decimal(0))
        nuovo_ricavi = sum((r.importi[p] for r in risultato.righe if r.sezione == "Ricavi Operativi"), Decimal(0))
        v = vecchi.get(p)
        out.append({
            "periodo": _periodo_txt(*p),
            "anno": p[0], "mese": p[1],
            "gia_presente": v is not None,
            "netto_vecchio": round(float(v["netto"]), 2) if v else None,
            "netto_nuovo": round(float(nuovo_netto), 2),
            "differenza": round(float(nuovo_netto - v["netto"]), 2) if v else None,
            "ricavi_vecchi": round(float(v["ricavi"]), 2) if v else None,
            "ricavi_nuovi": round(float(nuovo_ricavi), 2),
        })
    return out


def _parse_o_400(xlsx: bytes, csv: Optional[bytes]):
    try:
        return parse_export(xlsx, csv)
    except ErroreImportUsali as e:
        raise HTTPException(400, str(e))


# ── Anagrafiche ──────────────────────────────────────────────────────────────

@router.get("/strutture")
def lista_strutture(db: Session = Depends(get_db), _=Depends(richiedi_utente_attivo)):
    validi = _strutture_valide(db)
    con_dati = {
        s for (s,) in db.query(UsaliRiga.struttura_code).distinct().all()
    }
    mappe = defaultdict(list)
    for m in db.query(UsaliStrutturaGestionale).all():
        mappe[m.struttura_code].append(m.id_gestionale)
    return [
        {"code": c, "nome": n, "ha_dati": c in con_dati, "id_gestionale": mappe.get(c, [])}
        for c, n in validi.items()
    ]


@router.get("/segmenti")
def lista_segmenti(struttura: str, db: Session = Depends(get_db), _=Depends(richiedi_utente_attivo)):
    segs = db.query(UsaliSegmento).filter_by(struttura_code=struttura).order_by(UsaliSegmento.id).all()
    return [{"id": s.id, "nome": s.nome} for s in segs]


@router.get("/periodi")
def periodi_disponibili(struttura: str, db: Session = Depends(get_db), _=Depends(richiedi_utente_attivo)):
    righe = (
        db.query(UsaliRiga.anno, UsaliRiga.mese).filter(UsaliRiga.struttura_code == struttura)
        .distinct().order_by(UsaliRiga.anno, UsaliRiga.mese).all()
    )
    anni = defaultdict(list)
    for a, m in righe:
        anni[a].append(m)
    return [{"anno": a, "mesi": ms} for a, ms in sorted(anni.items(), reverse=True)]


# ── Import ───────────────────────────────────────────────────────────────────

@router.post("/anteprima")
async def anteprima_import(
    xlsx: Optional[UploadFile] = File(None),
    csv: Optional[UploadFile] = File(None),
    struttura_code: Optional[str] = Form(None),
    db: Session = Depends(get_db),
    _=Depends(richiedi_admin),
):
    xlsx_bytes = await xlsx.read() if xlsx else b""
    csv_bytes = await csv.read() if csv else None
    if not xlsx_bytes:
        raise HTTPException(
            400, "Serve il file xlsx: senza la gerarchia (rientri) dell'xlsx non si possono "
                 "classificare i conti. Il CSV è solo facoltativo."
        )
    risultato = _parse_o_400(xlsx_bytes, csv_bytes or None)

    rilevata = _struttura_da_gestionale(db, risultato.id_struttura_gestionale)
    avvisi = list(risultato.avvisi)
    if risultato.id_struttura_gestionale is not None and rilevata is None:
        avvisi.append(
            f"Struttura gestionale {risultato.id_struttura_gestionale} non mappata: sceglierla a mano."
        )
    # Struttura non determinabile → anteprima comunque restituita: l'utente la sceglie e
    # rilancia l'anteprima (serve per il confronto con i mesi già presenti).
    struttura = (
        _risolvi_struttura(db, rilevata, struttura_code) if (rilevata or struttura_code) else None
    )

    sha = hashlib.sha256(xlsx_bytes).hexdigest()
    duplicato = _import_duplicato(db, struttura, sha) if struttura else None

    _pulisci_anteprime_scadute()
    token = uuid.uuid4().hex
    d = _storage_dir() / "_anteprima" / token
    d.mkdir(parents=True, exist_ok=True)
    (d / "file.xlsx").write_bytes(xlsx_bytes)
    if csv_bytes:
        (d / "file.csv").write_bytes(csv_bytes)
    (d / "meta.json").write_text(json.dumps({
        "nome_xlsx": _nome_sicuro(xlsx.filename),
        "nome_csv": _nome_sicuro(csv.filename) if csv_bytes else None,
        "sha256": sha,
        "sha256_csv": hashlib.sha256(csv_bytes).hexdigest() if csv_bytes else None,
        "struttura_rilevata": rilevata,
    }))

    return {
        "token": token,
        "nome_xlsx": xlsx.filename,
        "nome_csv": csv.filename if csv_bytes else None,
        "struttura_rilevata": rilevata,
        "id_struttura_gestionale": risultato.id_struttura_gestionale,
        "struttura": struttura,
        "quadratura_ok": risultato.quadratura_ok,
        "csv_arricchito": risultato.csv_arricchito,
        "n_conti": len(risultato.righe),
        "periodi": [_periodo_txt(*p) for p in risultato.periodi],
        "confronto": _confronto_mesi(db, struttura, risultato) if struttura else [],
        "gia_importato": (
            {"id": duplicato.id, "data": duplicato.created_at.isoformat() if duplicato.created_at else None}
            if duplicato else None
        ),
        "avvisi": avvisi,
    }


@router.post("/conferma")
def conferma_import(
    body: dict = Body(...),
    db: Session = Depends(get_db),
    utente: User = Depends(richiedi_admin),
):
    d, meta, xlsx_bytes, csv_bytes = _leggi_anteprima(body.get("token"))
    risultato = _parse_o_400(xlsx_bytes, csv_bytes)
    struttura = _risolvi_struttura(db, meta.get("struttura_rilevata"), body.get("struttura_code"))

    duplicato = _import_duplicato(db, struttura, meta["sha256"])
    if duplicato:
        raise HTTPException(
            409, f"Questo file è già stato importato (import #{duplicato.id}): nessuna modifica."
        )

    periodi = risultato.periodi
    imp = UsaliImport(
        struttura_code=struttura,
        nome_file_xlsx=meta["nome_xlsx"], nome_file_csv=meta.get("nome_csv"),
        sha256=meta["sha256"], sha256_csv=meta.get("sha256_csv"),
        utente_id=getattr(utente, "id", None),
        periodi=[_periodo_txt(*p) for p in periodi],
        n_righe=len(risultato.righe) * len(periodi),
        quadratura_ok=risultato.quadratura_ok,
        csv_arricchito=risultato.csv_arricchito,
        avvisi=risultato.avvisi or None,
        note=(body.get("note") or None),
    )
    copiati = []
    try:
        db.add(imp)
        db.flush()

        filtro = or_(*[and_(UsaliRiga.anno == a, UsaliRiga.mese == m) for a, m in periodi])
        n_sostituite = (
            db.query(UsaliRiga).filter(UsaliRiga.struttura_code == struttura, filtro)
            .delete(synchronize_session=False)
        )
        db.bulk_insert_mappings(UsaliRiga, [
            {
                "import_id": imp.id, "struttura_code": struttura, "anno": a, "mese": m,
                "ordine": r.ordine, "percorso": r.percorso, "percorso_testo": r.percorso_testo,
                "sezione": r.sezione, "reparto": r.reparto, "sottoreparto": r.sottoreparto,
                "conto_nome": r.conto_nome, "id_conti": r.id_conti, "id_reparto": r.id_reparto,
                "tipo": r.tipo, "ammortamento": r.ammortamento, "importo": r.importi[(a, m)],
            }
            for r in risultato.righe for (a, m) in periodi
        ])

        dest = _storage_dir() / struttura / str(periodi[-1][0])
        dest.mkdir(parents=True, exist_ok=True)
        p_xlsx = dest / f"{imp.id}_{meta['nome_xlsx']}"
        shutil.copyfile(d / "file.xlsx", p_xlsx)
        copiati.append(p_xlsx)
        imp.percorso_xlsx = str(p_xlsx)
        if csv_bytes:
            p_csv = dest / f"{imp.id}_{meta['nome_csv']}"
            shutil.copyfile(d / "file.csv", p_csv)
            copiati.append(p_csv)
            imp.percorso_csv = str(p_csv)
        db.commit()
    except Exception:
        db.rollback()
        for p in copiati:
            p.unlink(missing_ok=True)
        raise
    shutil.rmtree(d, ignore_errors=True)

    return {
        "import_id": imp.id, "struttura": struttura,
        "periodi": imp.periodi, "righe_inserite": imp.n_righe, "righe_sostituite": n_sostituite,
    }


@router.get("/import")
def storico_import(
    struttura: Optional[str] = None,
    db: Session = Depends(get_db),
    _=Depends(richiedi_utente_attivo),
):
    q = db.query(UsaliImport, User.username).outerjoin(User, User.id == UsaliImport.utente_id)
    if struttura:
        q = q.filter(UsaliImport.struttura_code == struttura)
    # Mesi per cui ciascun import è ancora quello "in vigore"
    vigenti = defaultdict(list)
    for imp_id, a, m in db.query(UsaliRiga.import_id, UsaliRiga.anno, UsaliRiga.mese).distinct().all():
        vigenti[imp_id].append(_periodo_txt(a, m))
    return [
        {
            "id": i.id, "struttura": i.struttura_code,
            "nome_file_xlsx": i.nome_file_xlsx, "nome_file_csv": i.nome_file_csv,
            "ha_xlsx": bool(i.percorso_xlsx), "ha_csv": bool(i.percorso_csv),
            "utente": username,
            "created_at": i.created_at.isoformat() if i.created_at else None,
            "periodi": i.periodi, "periodi_in_vigore": sorted(vigenti.get(i.id, [])),
            "n_righe": i.n_righe, "quadratura_ok": i.quadratura_ok,
            "csv_arricchito": i.csv_arricchito, "avvisi": i.avvisi or [], "note": i.note,
        }
        for i, username in q.order_by(UsaliImport.id.desc()).all()
    ]


@router.delete("/import/{import_id}")
def elimina_import(
    import_id: int,
    conferma: bool = False,
    db: Session = Depends(get_db),
    _=Depends(richiedi_admin),
):
    """Elimina un import: le sue righe ancora in vigore, la voce di storico e i file originali.
    I mesi tolti tornano "non importati": un import precedente che quei mesi li aveva già
    sostituiti NON viene ripristinato (le sue righe non esistono più) — va reimportato."""
    imp = db.get(UsaliImport, import_id)
    if not imp:
        raise HTTPException(404, "Import non trovato")
    if not conferma:
        raise HTTPException(400, "Eliminazione non confermata: aggiungere ?conferma=true")

    periodi = sorted({
        _periodo_txt(a, m) for a, m in
        db.query(UsaliRiga.anno, UsaliRiga.mese).filter(UsaliRiga.import_id == import_id).distinct()
    })
    n_righe = db.query(UsaliRiga).filter(UsaliRiga.import_id == import_id).delete(synchronize_session=False)
    file_da_cancellare = [p for p in (imp.percorso_xlsx, imp.percorso_csv) if p]
    struttura = imp.struttura_code
    db.delete(imp)
    db.commit()
    for p in file_da_cancellare:
        Path(p).unlink(missing_ok=True)
    return {"import_id": import_id, "struttura": struttura,
            "righe_eliminate": n_righe, "periodi_rimossi": periodi}


@router.get("/import/{import_id}/file/{tipo}")
def scarica_file_originale(
    import_id: int, tipo: str,
    db: Session = Depends(get_db),
    _=Depends(richiedi_utente_attivo),
):
    imp = db.get(UsaliImport, import_id)
    if not imp or tipo not in ("xlsx", "csv"):
        raise HTTPException(404, "Import non trovato")
    percorso = imp.percorso_xlsx if tipo == "xlsx" else imp.percorso_csv
    nome = imp.nome_file_xlsx if tipo == "xlsx" else imp.nome_file_csv
    if not percorso or not Path(percorso).exists():
        raise HTTPException(404, "File originale non disponibile")
    return FileResponse(percorso, filename=nome)


# ── Report ───────────────────────────────────────────────────────────────────

def _carica_righe(db: Session, struttura: str, anno: int, mese_da: int, mese_a: int,
                  segmento_id: Optional[int]) -> dict:
    """Foglie di (anno, mesi ≤ mese_a) e (anno-1, mese_da..mese_a), già filtrate sul segmento
    se richiesto. Unica lettura usata da report, export e KPI."""
    if not (1 <= mese_da <= mese_a <= 12):
        raise HTTPException(400, "Intervallo mesi non valido")
    righe = (
        db.query(UsaliRiga)
        .filter(
            UsaliRiga.struttura_code == struttura,
            or_(
                and_(UsaliRiga.anno == anno, UsaliRiga.mese <= mese_a),
                and_(UsaliRiga.anno == anno - 1, UsaliRiga.mese.between(mese_da, mese_a)),
            ),
        )
        .all()
    )
    mesi_importati = {(r.anno, r.mese) for r in righe}

    errori_segmento = []
    segmento = None
    if segmento_id is not None:
        segmento = db.get(UsaliSegmento, segmento_id)
        if not segmento or segmento.struttura_code != struttura:
            raise HTTPException(404, "Segmento non trovato per questa struttura")
        regole = db.query(UsaliRegolaSegmento).filter_by(struttura_code=struttura).all()
        quote, errori = assegna_segmenti(righe, regole)
        righe = [
            {
                "anno": q.riga.anno, "mese": q.riga.mese, "importo": q.importo,
                "sezione": q.riga.sezione, "reparto": q.riga.reparto,
                "sottoreparto": q.riga.sottoreparto, "conto_nome": q.riga.conto_nome,
                "ordine": q.riga.ordine, "percorso_testo": q.riga.percorso_testo,
                "tipo": q.riga.tipo,
            }
            for q in quote if q.segmento_id == segmento_id
        ]
        visti = {}
        for e in errori:
            k = (e.riga.percorso_testo, e.motivo)
            visti.setdefault(k, []).append(_periodo_txt(e.riga.anno, e.riga.mese))
        errori_segmento = [
            {"conto": k[0], "motivo": k[1], "periodi": sorted(v)} for k, v in visti.items()
        ]

    return {
        "righe": righe,
        "mesi_importati": mesi_importati,
        "segmento": {"id": segmento.id, "nome": segmento.nome} if segmento else None,
        "errori_segmento": errori_segmento,
    }


def _calcola_report(db: Session, struttura: str, anno: int, mese_da: int, mese_a: int,
                    segmento_id: Optional[int]) -> dict:
    dati = _carica_righe(db, struttura, anno, mese_da, mese_a, segmento_id)
    prec_disponibile = any(a == anno - 1 for a, _ in dati["mesi_importati"])
    rep = costruisci_report(dati["righe"], anno, mese_da, mese_a, dati["mesi_importati"], prec_disponibile)
    rep.update({
        "struttura": struttura,
        "segmento": dati["segmento"],
        "errori_segmento": dati["errori_segmento"],
    })
    return rep


@router.get("/report")
def report_conto_economico(
    struttura: str,
    anno: int,
    mese_da: int = Query(1, ge=1, le=12),
    mese_a: int = Query(12, ge=1, le=12),
    segmento_id: Optional[int] = None,
    db: Session = Depends(get_db),
    _=Depends(richiedi_utente_attivo),
):
    return _calcola_report(db, struttura, anno, mese_da, mese_a, segmento_id)


_BLU = "1E3A5F"
_FMT_EURO = '#,##0.00 "€";[Red]-#,##0.00 "€"'
_FMT_PERC = '0.0"%";[Red]-0.0"%"'
_FILL_TIPO = {
    "sezione": PatternFill("solid", fgColor="E2E8F0"),
    "subtotale": PatternFill("solid", fgColor="DBEAFE"),
    "totale": PatternFill("solid", fgColor="1E3A5F"),
}


@router.get("/report/export")
def export_conto_economico(
    struttura: str,
    anno: int,
    mese_da: int = Query(1, ge=1, le=12),
    mese_a: int = Query(12, ge=1, le=12),
    segmento_id: Optional[int] = None,
    db: Session = Depends(get_db),
    _=Depends(richiedi_utente_attivo),
):
    rep = _calcola_report(db, struttura, anno, mese_da, mese_a, segmento_id)
    nomi_mesi = [MESI_IT[m - 1].capitalize() for m in rep["mesi"]]
    intest = ["Voce"] + nomi_mesi + ["Totale periodo", "% ricavi", f"Progressivo {anno}"]
    if rep["prec_disponibile"]:
        intest += [f"Stesso periodo {anno - 1}", f"Δ vs {anno - 1}"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Conto Economico"
    titolo = f"Conto Economico USALI — {struttura}"
    if rep["segmento"]:
        titolo += f" ({rep['segmento']['nome']})"
    titolo += f" — {nomi_mesi[0]}–{nomi_mesi[-1]} {anno}"
    ws.append([titolo])
    ws.cell(1, 1).font = Font(bold=True, size=13)
    ws.append([])
    ws.append(intest)
    for c in range(1, len(intest) + 1):
        cella = ws.cell(3, c)
        cella.font = Font(bold=True, color="FFFFFF")
        cella.fill = PatternFill("solid", fgColor=_BLU)
        cella.alignment = Alignment(horizontal="center", wrap_text=True)

    for r in rep["righe"]:
        valori = [r["nome"]] + r["valori"] + [r["totale"], r["perc_ricavi"], r["progressivo"]]
        if rep["prec_disponibile"]:
            valori += [r["anno_prec"], r["delta_anno_prec"]]
        ws.append(valori)
        n = ws.max_row
        ws.cell(n, 1).alignment = Alignment(indent=r["livello"] * 2)
        for c in range(2, len(valori) + 1):
            ws.cell(n, c).number_format = _FMT_PERC if intest[c - 1] == "% ricavi" else _FMT_EURO
        fill = _FILL_TIPO.get(r["tipo"])
        if fill:
            for c in range(1, len(valori) + 1):
                ws.cell(n, c).fill = fill
                ws.cell(n, c).font = Font(bold=True, color="FFFFFF" if r["tipo"] == "totale" else "000000")

    if rep["errori_segmento"]:
        ws.append([])
        ws.append([f"⚠ {len(rep['errori_segmento'])} conti non assegnati a nessun segmento: esclusi dai totali"])

    ws.column_dimensions["A"].width = 55
    for c in range(2, len(intest) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 15
    ws.freeze_panes = "B4"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nome = f"CE_USALI_{struttura}_{anno}_{rep['mese_da']:02d}-{rep['mese_a']:02d}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nome}"'},
    )


# ── KPI ──────────────────────────────────────────────────────────────────────

@router.get("/kpi")
def kpi_conto_economico(
    struttura: str,
    anno: int,
    mese_da: int = Query(1, ge=1, le=12),
    mese_a: int = Query(12, ge=1, le=12),
    segmento_id: Optional[int] = None,
    db: Session = Depends(get_db),
    _=Depends(richiedi_utente_attivo),
):
    """KPI mensili e di periodo. Con un segmento, gli importi sono quelli del segmento ma le
    camere restano quelle dell'intera struttura (la produzione non è segmentata)."""
    dati = _carica_righe(db, struttura, anno, mese_da, mese_a, segmento_id)
    out = calcola_kpi(
        dati["righe"], carica_mappa_categorie(db), produzione_mensile(db, struttura, anno),
        anno, mese_da, mese_a, dati["mesi_importati"],
    )
    out.update({"struttura": struttura, "segmento": dati["segmento"],
                "errori_segmento": dati["errori_segmento"]})
    return out


# ── Categorie KPI ────────────────────────────────────────────────────────────

def _valida_categoria(body: dict) -> dict:
    conto = (body.get("conto_nome") or "").strip()
    categoria = (body.get("categoria") or "").strip()
    if not conto:
        raise HTTPException(400, "Nome conto obbligatorio")
    if categoria not in CATEGORIE_KPI:
        raise HTTPException(400, f"Categoria non valida: usare una tra {', '.join(CATEGORIE_KPI)}")
    return {"conto_nome": conto, "reparto": (body.get("reparto") or "").strip(), "categoria": categoria}


def _fmt_categoria(c: UsaliCategoriaKpi) -> dict:
    return {"id": c.id, "conto_nome": c.conto_nome, "reparto": c.reparto, "categoria": c.categoria}


@router.get("/categorie-kpi")
def lista_categorie_kpi(db: Session = Depends(get_db), _=Depends(richiedi_utente_attivo)):
    righe = db.query(UsaliCategoriaKpi).order_by(UsaliCategoriaKpi.categoria, UsaliCategoriaKpi.conto_nome).all()
    return {"categorie": list(CATEGORIE_KPI), "mappature": [_fmt_categoria(c) for c in righe]}


def _esiste_categoria(db: Session, conto: str, reparto: str, escludi_id: Optional[int] = None) -> bool:
    q = db.query(UsaliCategoriaKpi).filter(
        func.lower(UsaliCategoriaKpi.conto_nome) == conto.lower(),
        func.lower(UsaliCategoriaKpi.reparto) == reparto.lower(),
    )
    if escludi_id:
        q = q.filter(UsaliCategoriaKpi.id != escludi_id)
    return db.query(q.exists()).scalar()


@router.post("/categorie-kpi")
def crea_categoria_kpi(body: dict = Body(...), db: Session = Depends(get_db), _=Depends(richiedi_admin)):
    dati = _valida_categoria(body)
    if _esiste_categoria(db, dati["conto_nome"], dati["reparto"]):
        raise HTTPException(409, "Esiste già una mappatura per questo conto/reparto")
    c = UsaliCategoriaKpi(**dati)
    db.add(c)
    db.commit()
    return _fmt_categoria(c)


@router.put("/categorie-kpi/{cat_id}")
def aggiorna_categoria_kpi(cat_id: int, body: dict = Body(...), db: Session = Depends(get_db),
                           _=Depends(richiedi_admin)):
    c = db.get(UsaliCategoriaKpi, cat_id)
    if not c:
        raise HTTPException(404, "Mappatura non trovata")
    dati = _valida_categoria({**_fmt_categoria(c), **body})
    if _esiste_categoria(db, dati["conto_nome"], dati["reparto"], escludi_id=cat_id):
        raise HTTPException(409, "Esiste già una mappatura per questo conto/reparto")
    for k, v in dati.items():
        setattr(c, k, v)
    db.commit()
    return _fmt_categoria(c)


@router.delete("/categorie-kpi/{cat_id}")
def elimina_categoria_kpi(cat_id: int, db: Session = Depends(get_db), _=Depends(richiedi_admin)):
    c = db.get(UsaliCategoriaKpi, cat_id)
    if not c:
        raise HTTPException(404, "Mappatura non trovata")
    db.delete(c)
    db.commit()
    return {"ok": True}


@router.get("/categorie-kpi/non-categorizzati")
def conti_non_categorizzati(
    struttura: Optional[str] = None,
    db: Session = Depends(get_db),
    _=Depends(richiedi_utente_attivo),
):
    """Conti di costo presenti negli import senza categoria KPI (i conti nuovi che arrivano dal
    gestionale finiscono qui). Assegnare 'altro' per toglierli dall'elenco."""
    q = db.query(
        UsaliRiga.conto_nome, UsaliRiga.reparto, UsaliRiga.sezione, UsaliRiga.tipo,
        func.sum(UsaliRiga.importo), func.array_agg(func.distinct(UsaliRiga.struttura_code)),
    ).filter(UsaliRiga.tipo == "costo")
    if struttura:
        q = q.filter(UsaliRiga.struttura_code == struttura)
    gruppi = q.group_by(UsaliRiga.conto_nome, UsaliRiga.reparto, UsaliRiga.sezione, UsaliRiga.tipo).all()

    mappa = carica_mappa_categorie(db)
    out = []
    for conto, reparto, sezione, tipo, totale, strutture in gruppi:
        riga = {"conto_nome": conto, "reparto": reparto, "tipo": tipo}
        if categoria_di(riga, mappa) is None:
            out.append({
                "conto_nome": conto, "reparto": reparto, "sezione": sezione,
                "totale": round(float(totale or 0), 2), "strutture": sorted(strutture or []),
            })
    out.sort(key=lambda x: x["totale"])   # costi più pesanti (più negativi) in cima
    return out
