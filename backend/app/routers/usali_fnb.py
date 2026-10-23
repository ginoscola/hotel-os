"""USALI → Food & Beverage: analisi per punto vendita (logica in services/usali_fnb.py).

  GET /usali/fnb/report?struttura=DPH|CLB|INT|BON|GRUPPO&anno=&mese_da=&mese_a=
  GET /usali/fnb/materia-prima?struttura=&anno=&mese_da=&mese_a=
  GET /usali/fnb/costo-pasto?anno=&mese_da=&mese_a=   → costo per colazione / pranzo / cena
  GET /usali/fnb/coperti?anno=          → coperti MMS (Menu) e scontrini BON per mese
  PUT /usali/fnb/coperti                → {struttura_code, anno, mese, valore|null} (admin)
  GET|PUT /usali/fnb/parametri          → peso cena/pranzo + centro di costo → punto vendita
  GET /usali/fnb/report/export          → xlsx del riepilogo (stessi parametri di /report)
"""
import io

import openpyxl
from fastapi import APIRouter, Body, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from app.auth import richiedi_admin, richiedi_utente_attivo
from app.database import get_db
from app.models.usali import UsaliFnbCoperti
from app.services import usali_fnb as fnb
from app.utils.locale_it import MESI_IT

router = APIRouter(prefix="/usali/fnb", tags=["usali-fnb"])


def _verifica(struttura: str, mese_da: int, mese_a: int, gruppo: bool = True):
    valide = fnb.STRUTTURE + (['GRUPPO'] if gruppo else [])
    if struttura not in valide:
        raise HTTPException(400, f"Struttura non valida: {struttura}")
    if not (1 <= mese_da <= mese_a <= 12):
        raise HTTPException(400, "Intervallo mesi non valido")


def calcola_report(db: Session, struttura: str, anno: int, mese_da: int, mese_a: int) -> dict:
    parametri = fnb.leggi_parametri(db)
    peso = parametri['peso_cena']
    if struttura == 'GRUPPO':
        reports = [fnb.costruisci_report(fnb.carica_dati(db, s, anno, parametri), mese_da, mese_a, peso)
                   for s in fnb.STRUTTURE]
        out = fnb.report_gruppo(reports, mese_da, mese_a)
    else:
        out = fnb.costruisci_report(fnb.carica_dati(db, struttura, anno, parametri), mese_da, mese_a, peso)
    out.update({'anno': anno, 'mese_da': mese_da, 'mese_a': mese_a, 'peso_cena': peso})
    return out


@router.get("/report")
def report_fnb(
    struttura: str, anno: int,
    mese_da: int = Query(1, ge=1, le=12), mese_a: int = Query(12, ge=1, le=12),
    db: Session = Depends(get_db), _=Depends(richiedi_utente_attivo),
):
    _verifica(struttura, mese_da, mese_a)
    return calcola_report(db, struttura, anno, mese_da, mese_a)


# (chiave, intestazione, formato) — stesso ordine delle colonne a schermo (pages/usaliFnb/colonne.jsx)
_FMT_EURO, _FMT_PERC, _FMT_INT = '#,##0.00 "€"', '0.0"%"', '#,##0'
COLONNE_EXPORT = [
    ('ricavi', 'Ricavi', _FMT_EURO), ('ricavi_contabili', 'Ricavi contabilità', _FMT_EURO),
    ('food', 'Food', _FMT_EURO), ('beverage', 'Beverage', _FMT_EURO),
    ('materia_prima_perc', 'Materia prima %', _FMT_PERC), ('personale', 'Personale', _FMT_EURO),
    ('prime_cost_perc', 'Prime cost %', _FMT_PERC), ('altri_costi', 'Altri costi', _FMT_EURO),
    ('margine', 'Margine', _FMT_EURO), ('margine_perc', 'Margine %', _FMT_PERC),
    ('coperti', 'Coperti', _FMT_INT), ('ricavo_coperto', 'Ricavo / coperto', _FMT_EURO),
    ('mp_coperto', 'Materia prima / coperto', _FMT_EURO), ('costo_coperto', 'Costo pieno / coperto', _FMT_EURO),
    ('adesione_perc', 'Adesione (pasti / presenze)', _FMT_PERC),
]


@router.get("/report/export")
def export_report_fnb(
    struttura: str, anno: int,
    mese_da: int = Query(1, ge=1, le=12), mese_a: int = Query(12, ge=1, le=12),
    db: Session = Depends(get_db), _=Depends(richiedi_utente_attivo),
):
    _verifica(struttura, mese_da, mese_a)
    rep = calcola_report(db, struttura, anno, mese_da, mese_a)
    nome_s = 'Gruppo' if struttura == 'GRUPPO' else fnb.NOME_STRUTTURA[struttura]
    periodo = f"{MESI_IT[mese_da - 1].capitalize()}–{MESI_IT[mese_a - 1].capitalize()} {anno}"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Food & Beverage"
    ws.append([f"Food & Beverage — {nome_s} — {periodo}"])
    ws.cell(1, 1).font = Font(bold=True, size=13)
    ws.append([])
    intest = ['Punto vendita'] + [c[1] for c in COLONNE_EXPORT]
    ws.append(intest)
    for c in range(1, len(intest) + 1):
        cella = ws.cell(3, c)
        cella.font = Font(bold=True, color="FFFFFF")
        cella.fill = PatternFill("solid", fgColor="1E3A5F")
        cella.alignment = Alignment(horizontal="center", wrap_text=True)
    fill_tot = PatternFill("solid", fgColor="1E3A5F")
    fill_punto = PatternFill("solid", fgColor="F1F5F9")
    for r in rep['righe']:
        label = r['label'] + (' (costi stimati)' if r['stima_costi'] else '')
        ws.append([label] + [r['periodo'][k] for k, _, _ in COLONNE_EXPORT])
        n = ws.max_row
        ws.cell(n, 1).alignment = Alignment(indent=2 if r['tipo'] == 'sub' else 0)
        for i, (_, _, fmt) in enumerate(COLONNE_EXPORT, start=2):
            ws.cell(n, i).number_format = fmt
        if r['tipo'] in ('totale', 'punto'):
            for c in range(1, len(intest) + 1):
                ws.cell(n, c).fill = fill_tot if r['tipo'] == 'totale' else fill_punto
                ws.cell(n, c).font = Font(bold=True, color="FFFFFF" if r['tipo'] == 'totale' else "000000")
    ws.append([])
    ws.append([f"Costi = acquisti registrati in contabilità; costi di pranzo/cena e ospiti/esterni stimati "
               f"con una cena = {rep['peso_cena']} pranzi. Personale dai cedolini. Rapporti solo sui mesi "
               "con contabilità importata."])
    for a in rep['avvisi']:
        ws.append([f"⚠ {a}"])
    ws.column_dimensions['A'].width = 34
    for c in range(2, len(intest) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 14
    ws.freeze_panes = 'B4'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nome = f"FnB_{struttura}_{anno}_{mese_da:02d}-{mese_a:02d}.xlsx"
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nome}"'},
    )


@router.get("/costo-pasto")
def costo_pasto_fnb(
    anno: int,
    mese_da: int = Query(1, ge=1, le=12), mese_a: int = Query(12, ge=1, le=12),
    db: Session = Depends(get_db), _=Depends(richiedi_utente_attivo),
):
    """Costo per colazione (DPH/CLB/INT) e per pasto del ristorante (pranzo/cena CLB e INT,
    Mare Mosso ospiti/esterni), dalle stesse righe del riepilogo."""
    _verifica('GRUPPO', mese_da, mese_a)
    parametri = fnb.leggi_parametri(db)
    reports = {s: fnb.costruisci_report(fnb.carica_dati(db, s, anno, parametri), mese_da, mese_a,
                                        parametri['peso_cena'])
               for s in ('DPH', 'CLB', 'INT')}
    return {'anno': anno, 'mese_da': mese_da, 'mese_a': mese_a, 'peso_cena': parametri['peso_cena'],
            'righe': fnb.costo_pasto(reports)}


@router.get("/materia-prima")
def materia_prima_fnb(
    struttura: str, anno: int,
    mese_da: int = Query(1, ge=1, le=12), mese_a: int = Query(12, ge=1, le=12),
    db: Session = Depends(get_db), _=Depends(richiedi_utente_attivo),
):
    _verifica(struttura, mese_da, mese_a)
    parametri = fnb.leggi_parametri(db)
    strutture = fnb.STRUTTURE if struttura == 'GRUPPO' else [struttura]
    out = {'conti': [], 'non_categorizzati': [], 'totale_non_categorizzati': 0.0}
    for s in strutture:
        r = fnb.materia_prima(fnb.carica_dati(db, s, anno, parametri), mese_da, mese_a)
        out['conti'] += r['conti']
        out['non_categorizzati'] += r['non_categorizzati']
        out['totale_non_categorizzati'] = round(out['totale_non_categorizzati'] + r['totale_non_categorizzati'], 2)
    out['conti'].sort(key=lambda r: -r['totale'])
    out['non_categorizzati'].sort(key=lambda r: -r['totale'])
    return out


@router.get("/coperti")
def lista_coperti(anno: int, db: Session = Depends(get_db), _=Depends(richiedi_utente_attivo)):
    """Coperti manuali per mese, con i pasti degli ospiti Du Parc (da Produzione) per confronto:
    clienti esterni Mare Mosso = coperti Menu − pasti ospiti."""
    manuali = fnb.carica_coperti_manuali(db, anno)
    _, pasti = fnb._carica_produzione(db, 'DPH', anno)
    mesi = []
    for m in range(1, 13):
        ospiti = pasti.get(m, {}).get('pranzo', 0) + pasti.get(m, {}).get('cena', 0)
        menu = manuali.get(('MMS', m))
        mesi.append({
            'mese': m, 'mms_coperti': menu, 'bon_scontrini': manuali.get(('BON', m)),
            'ospiti_dph': ospiti,
            'mms_esterni': None if menu is None else menu - ospiti,
        })
    return {'anno': anno, 'mesi': mesi}


@router.put("/coperti")
def salva_coperti(body: dict = Body(...), db: Session = Depends(get_db), _=Depends(richiedi_admin)):
    struttura = body.get('struttura_code')
    if struttura not in ('MMS', 'BON'):
        raise HTTPException(400, "struttura_code deve essere MMS o BON")
    try:
        anno, mese = int(body['anno']), int(body['mese'])
    except (KeyError, TypeError, ValueError):
        raise HTTPException(400, "anno e mese obbligatori")
    if not 1 <= mese <= 12:
        raise HTTPException(400, "Mese non valido")
    valore = body.get('valore')
    riga = db.query(UsaliFnbCoperti).filter_by(struttura_code=struttura, anno=anno, mese=mese).first()
    if valore is None or valore == '':
        if riga:
            db.delete(riga)
            db.commit()
        return {'struttura_code': struttura, 'anno': anno, 'mese': mese, 'valore': None}
    try:
        valore = int(valore)
    except (TypeError, ValueError):
        raise HTTPException(400, "Il valore deve essere un numero intero")
    if valore < 0:
        raise HTTPException(400, "Il valore non può essere negativo")
    if riga:
        riga.valore = valore
    else:
        db.add(UsaliFnbCoperti(struttura_code=struttura, anno=anno, mese=mese, valore=valore))
    db.commit()
    return {'struttura_code': struttura, 'anno': anno, 'mese': mese, 'valore': valore}


def _parametri_out(db: Session) -> dict:
    p = fnb.leggi_parametri(db)
    centri = [{**c, 'punto': p['cc_punto'].get(c['code']) or fnb.punto_default_cc(c['code']),
               'punto_default': fnb.punto_default_cc(c['code'])}
              for c in fnb.centri_fnb(db)]
    return {'peso_cena': p['peso_cena'], 'peso_cena_default': fnb.PESO_CENA_DEFAULT, 'centri': centri}


@router.get("/parametri")
def get_parametri(db: Session = Depends(get_db), _=Depends(richiedi_utente_attivo)):
    return _parametri_out(db)


@router.put("/parametri")
def put_parametri(body: dict = Body(...), db: Session = Depends(get_db), _=Depends(richiedi_admin)):
    try:
        fnb.salva_parametri(db, body)
    except (ValueError, TypeError) as e:
        raise HTTPException(400, str(e))
    return _parametri_out(db)
