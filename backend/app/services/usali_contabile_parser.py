"""Parser dell'export Flexmonster del gestionale contabile (Conto Economico già riclassificato USALI).

- xlsx (obbligatorio): gerarchia dal rientro della cella in colonna A (`alignment.indent`),
  importi con decimali pieni, mappa colonna → (anno, mese) da riga 2 (anno, forward-fill) e
  riga 3 (nome mese italiano). L'ultima riga 'Gran Totale' è il controllo di quadratura.
- csv (opzionale): stessa tabella, stesso ordine righe; importi arrotondati all'euro (ignorati),
  ma nelle foglie la colonna A è 'Nome|URL|Nome|URL' con parametri utili (tipo movimento,
  id conto, id reparto, ammortamento, struttura gestionale).

Funzioni pure (nessun accesso DB): `parse_export(xlsx_bytes, csv_bytes=None)`.
"""

import csv
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal, ROUND_HALF_UP
from typing import Optional
from urllib.parse import unquote_plus

import openpyxl

from app.utils.locale_it import MESI_IT

FOGLIO = "Flexmonster Pivot Table"
RIGA_ANNI = 2
RIGA_MESI = 3
PRIMA_RIGA_DATI = 4
GRAN_TOTALE = "Gran Totale"
TOLLERANZA_QUADRATURA = 0.01

# Ordine di ricerca dalla radice: il primo nodo del percorso che compare qui determina la
# sezione (Deprezzamenti contiene a sua volta 'Costi Generali', ma sta più in alto).
SEZIONI = ("Ricavi Operativi", "Costi Operativi", "Costi Generali", "Non Operativi", "Deprezzamenti")

CENT = Decimal("0.0001")   # 4 decimali: il gestionale ripartisce costi con frazioni di centesimo


class ErroreImportUsali(ValueError):
    """Errore bloccante: l'import non può procedere (messaggio mostrato all'utente)."""


@dataclass
class RigaConto:
    ordine: int                         # numero riga nel foglio xlsx
    percorso: list
    sezione: Optional[str]
    reparto: Optional[str]
    sottoreparto: Optional[str]
    conto_nome: str
    tipo: str                           # 'ricavo' | 'costo'
    ammortamento: bool
    importi: dict                       # (anno, mese) -> Decimal (4 decimali)
    id_conti: Optional[int] = None
    id_reparto: Optional[int] = None

    @property
    def percorso_testo(self) -> str:
        return " > ".join(self.percorso)


@dataclass
class RisultatoParsing:
    righe: list
    periodi: list                       # [(anno, mese), ...] in ordine di colonna
    quadratura_ok: bool
    differenze_quadratura: dict         # (anno, mese) -> differenza (solo se fuori tolleranza)
    csv_arricchito: bool = False
    id_struttura_gestionale: Optional[int] = None
    avvisi: list = field(default_factory=list)


def _dec(v) -> Decimal:
    if v is None or v == "":
        return Decimal("0.00")
    return Decimal(str(v)).quantize(CENT, rounding=ROUND_HALF_UP)


def _mese_da_nome(nome) -> Optional[int]:
    if not isinstance(nome, str):
        return None
    prefisso = nome.strip().lower()[:3]
    return MESI_IT.index(prefisso) + 1 if prefisso in MESI_IT else None


def _mappa_colonne(ws) -> dict:
    """colonna (1-based) → (anno, mese). Ignora le colonne totale."""
    mappa = {}
    anno_corrente = None
    for col in range(2, ws.max_column + 1):
        v_anno = ws.cell(RIGA_ANNI, col).value
        if v_anno not in (None, ""):
            try:
                anno_corrente = int(str(v_anno).strip())
            except ValueError:
                anno_corrente = None   # 'Totali' o simili: chiude il blocco anno
        mese = _mese_da_nome(ws.cell(RIGA_MESI, col).value)
        if anno_corrente is not None and mese is not None:
            mappa[col] = (anno_corrente, mese)
    return mappa


def _classifica(percorso: list) -> tuple:
    """(sezione, reparto, sottoreparto) dal percorso radice→foglia."""
    idx_sez = next((i for i, nodo in enumerate(percorso[:-1]) if nodo in SEZIONI), None)
    if idx_sez is None:
        return None, None, None
    sezione = percorso[idx_sez]
    intermedi = percorso[idx_sez + 1:-1]
    if not intermedi:
        return sezione, None, None
    reparto = intermedi[0]
    # Nodi che ripetono il nome del reparto (es. 'Altri Reparti Operativi' annidato 3 volte)
    # non aggiungono informazione.
    sotto = [n for n in intermedi[1:] if n != reparto]
    return sezione, reparto, (" > ".join(sotto) or None)


def _leggi_xlsx(xlsx_bytes: bytes):
    try:
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    except Exception as e:
        raise ErroreImportUsali(f"File xlsx non leggibile: {e}")
    if FOGLIO in wb.sheetnames:
        ws = wb[FOGLIO]
    elif len(wb.sheetnames) == 1:
        ws = wb.active
    else:
        raise ErroreImportUsali(f"Foglio '{FOGLIO}' non trovato nel file xlsx")
    return ws


def _parse_xlsx(xlsx_bytes: bytes) -> tuple:
    ws = _leggi_xlsx(xlsx_bytes)
    mappa = _mappa_colonne(ws)
    if not mappa:
        raise ErroreImportUsali("Nessuna colonna mese riconosciuta (righe 2-3: anno e mese in italiano)")

    # Righe dati fino a 'Gran Totale' (esclusa)
    nodi = []        # (riga, indent, nome)
    riga_totale = None
    for r in range(PRIMA_RIGA_DATI, ws.max_row + 1):
        cella = ws.cell(r, 1)
        nome = cella.value
        if nome is None or str(nome).strip() == "":
            continue
        nome = str(nome).strip()
        if nome == GRAN_TOTALE:
            riga_totale = r
            break
        nodi.append((r, int(cella.alignment.indent or 0), nome))
    if riga_totale is None:
        raise ErroreImportUsali("Riga 'Gran Totale' non trovata: il file non sembra un export completo")

    periodi = sorted(set(mappa.values()))
    righe = []
    subtotali = []   # (indice in nodi, riga, rientro, nome) delle intestazioni con importi
    stack = []   # [(indent, nome)]
    for i, (r, indent, nome) in enumerate(nodi):
        while stack and stack[-1][0] >= indent:
            stack.pop()
        stack.append((indent, nome))
        prossimo_indent = nodi[i + 1][1] if i + 1 < len(nodi) else -1
        is_foglia = prossimo_indent <= indent

        valori = {col: ws.cell(r, col).value for col in mappa}
        if not is_foglia:
            # Alcuni export (es. CLB) riportano i subtotali anche sulle intestazioni: non sono
            # conti, servono solo come controllo (devono coincidere con la somma delle foglie).
            if any(isinstance(v, (int, float)) and v != 0 for v in valori.values()):
                subtotali.append((i, r, indent, nome))
            continue

        percorso = [n for _, n in stack]
        sezione, reparto, sottoreparto = _classifica(percorso)
        importi = {p: Decimal("0.00") for p in periodi}
        for col, periodo in mappa.items():
            importi[periodo] += _dec(valori[col])
        righe.append(RigaConto(
            ordine=r, percorso=percorso, sezione=sezione, reparto=reparto,
            sottoreparto=sottoreparto, conto_nome=nome,
            tipo="ricavo" if sezione == "Ricavi Operativi" else "costo",
            ammortamento=(sezione == "Deprezzamenti"),
            importi=importi,
        ))

    # Quadratura: somma foglie (valori pieni, non arrotondati) vs Gran Totale, per colonna
    differenze = {}
    for col, periodo in mappa.items():
        somma = 0.0
        for riga in righe:
            v = ws.cell(riga.ordine, col).value
            somma += float(v) if isinstance(v, (int, float)) else 0.0
        v_tot = ws.cell(riga_totale, col).value
        totale = float(v_tot) if isinstance(v_tot, (int, float)) else 0.0
        if abs(somma - totale) > TOLLERANZA_QUADRATURA:
            differenze[periodo] = round(somma - totale, 2)

    _verifica_subtotali(ws, mappa, nodi, subtotali, {riga.ordine for riga in righe})

    nomi_per_riga = {r: nome for r, _, nome in nodi}
    nomi_per_riga[riga_totale] = GRAN_TOTALE
    return righe, periodi, differenze, nomi_per_riga


def _verifica_subtotali(ws, mappa: dict, nodi: list, subtotali: list, righe_foglia: set):
    """Ogni intestazione con importi deve valere la somma delle foglie sotto di lei (le righe
    successive con rientro maggiore). Se non torna, il file non è affidabile: errore."""
    for i, r, indent, nome in subtotali:
        discendenti = []
        for r2, ind2, _ in nodi[i + 1:]:
            if ind2 <= indent:
                break
            if r2 in righe_foglia:
                discendenti.append(r2)
        for col, (anno, mese) in mappa.items():
            v = ws.cell(r, col).value
            sub = float(v) if isinstance(v, (int, float)) else 0.0
            somma = 0.0
            for r2 in discendenti:
                x = ws.cell(r2, col).value
                somma += float(x) if isinstance(x, (int, float)) else 0.0
            if abs(somma - sub) > TOLLERANZA_QUADRATURA:
                raise ErroreImportUsali(
                    f"Riga {r} '{nome}', {MESI_IT[mese - 1]} {anno}: il subtotale del file "
                    f"({sub:,.2f} €) non coincide con la somma dei suoi conti ({somma:,.2f} €). "
                    f"Riesportare dal gestionale con tutte le righe espanse."
                )


_RE_ID_CONTI = re.compile(r"[?&]id_conti=(\d+)")
_RE_ID_REPARTO = re.compile(r"[?&]id_reparti_(\d+)=")
_RE_AZIONE = re.compile(r"[?&]Azione=([^&|]+)")
_RE_AMMORT = re.compile(r"[?&]Ammortamento=1(?:&|$)")
_RE_STRUTTURE = re.compile(r"[?&]id_strutture_concat_work=([^&|]+)")


def _parametri_url(cella: str) -> Optional[dict]:
    parti = cella.split("|")
    if len(parti) < 2:
        return None
    url = parti[1]
    m_conti = _RE_ID_CONTI.search(url)
    m_azione = _RE_AZIONE.search(url)
    m_rep = _RE_ID_REPARTO.search(url)
    m_str = _RE_STRUTTURE.search(url)
    strutture = []
    if m_str:
        strutture = [int(x) for x in unquote_plus(m_str.group(1)).split(",") if x.strip().isdigit()]
    return {
        "azione": m_azione.group(1) if m_azione else None,
        "id_conti": int(m_conti.group(1)) if m_conti else None,
        "id_reparto": int(m_rep.group(1)) if m_rep else None,
        "ammortamento": bool(_RE_AMMORT.search(url)),
        "strutture": [s for s in strutture if s != 0],
    }


def _arricchisci_da_csv(csv_bytes: bytes, righe: list, nomi_per_riga: dict, risultato: RisultatoParsing):
    """Abbina CSV↔xlsx per indice di riga. Qualunque incoerenza → arricchimento scartato
    per tutto il file (con avviso), l'import procede con i soli dati xlsx."""
    try:
        testo = csv_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        risultato.avvisi.append("CSV non in UTF-8: arricchimento (id conto/reparto) ignorato.")
        return
    righe_csv = list(csv.reader(io.StringIO(testo)))

    for r, nome_xlsx in nomi_per_riga.items():
        idx = r - 1
        if idx >= len(righe_csv) or not righe_csv[idx]:
            risultato.avvisi.append(
                f"CSV più corto dell'xlsx (riga {r} assente): arricchimento ignorato."
            )
            return
        nome_csv = righe_csv[idx][0].split("|", 1)[0].strip()
        if nome_csv != nome_xlsx:
            risultato.avvisi.append(
                f"Riga {r}: nome nel CSV '{nome_csv}' diverso dall'xlsx '{nome_xlsx}' — "
                f"i due file non corrispondono, arricchimento CSV ignorato."
            )
            return

    parametri = {}
    strutture = set()
    for riga in righe:
        p = _parametri_url(righe_csv[riga.ordine - 1][0])
        if p is None:
            continue
        parametri[riga.ordine] = p
        strutture.update(p["strutture"])

    for riga in righe:
        p = parametri.get(riga.ordine)
        if p is None:
            continue
        riga.id_conti = p["id_conti"]
        riga.id_reparto = p["id_reparto"]
        riga.ammortamento = p["ammortamento"]
        if p["azione"] == "ricerca_movimenti_attivi":
            riga.tipo = "ricavo"
        elif p["azione"] == "ricerca_movimenti_passivi":
            riga.tipo = "costo"

    risultato.csv_arricchito = True
    if len(strutture) == 1:
        risultato.id_struttura_gestionale = strutture.pop()
    elif len(strutture) > 1:
        risultato.avvisi.append(
            f"Il CSV cita più strutture del gestionale ({sorted(strutture)}): scegliere la struttura a mano."
        )


def _sembra_csv(dati: bytes) -> bool:
    return not dati.startswith(b"PK")


def parse_export(xlsx_bytes: bytes, csv_bytes: Optional[bytes] = None) -> RisultatoParsing:
    if not xlsx_bytes:
        raise ErroreImportUsali(
            "Serve il file xlsx: senza la gerarchia (rientri) dell'xlsx non si possono classificare i conti."
        )
    if _sembra_csv(xlsx_bytes):
        raise ErroreImportUsali(
            "Il file caricato come xlsx è un CSV. Serve l'export xlsx: il CSV non contiene la gerarchia "
            "dei conti né gli importi con i decimali."
        )

    righe, periodi, differenze, nomi_per_riga = _parse_xlsx(xlsx_bytes)
    if differenze:
        dettaglio = ", ".join(
            f"{MESI_IT[m - 1]} {a}: {d:+.2f} €" for (a, m), d in sorted(differenze.items())
        )
        raise ErroreImportUsali(
            f"Il file non quadra con il Gran Totale (somma conti − Gran Totale: {dettaglio}). "
            f"Riesportare dal gestionale."
        )
    risultato = RisultatoParsing(
        righe=righe, periodi=periodi,
        quadratura_ok=not differenze, differenze_quadratura=differenze,
    )
    for riga in righe:
        if riga.sezione is None:
            risultato.avvisi.append(
                f"Conto '{riga.percorso_testo}' fuori dalle sezioni USALI note: escluso dai subtotali."
            )
    if csv_bytes:
        _arricchisci_da_csv(csv_bytes, righe, nomi_per_riga, risultato)
    return risultato
