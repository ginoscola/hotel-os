"""Parser degli export Welcome di prenotazioni cancellate — due formati supportati.

**Formato "PrenotazioniWeb"** (CSV separato da virgola, prima riga "sep=," da scartare):
date ISO 8601, importi con punto decimale. Filtrato lato Welcome per mese di *prenotazione*.
Due colonne opzionali non native, leggibili se l'utente le aggiunge a mano al CSV: "Codice
Prenotazione"/"codicePrenotazione" e "Data cancellazione"/"dataCancellazione".

**Formato "Elenco Prenotazioni"** (CSV separato da **punto e virgola**, rilevato dalla presenza
di "Codice prenotazione" nell'intestazione): date italiane `gg/mm/aaaa`, importi con virgola
decimale, `Codice prenotazione` e `Data cancellazione` **nativi** (non opzionali). Filtrato lato
Welcome per mese di **arrivo**, non di prenotazione — la `Data prenotazione` in questo formato
spazia su mesi diversi, è normale.
⚠️ **Contiene versioni storiche della stessa prenotazione**: una modifica (importo, conferma,
ecc.) genera una nuova riga con lo stesso `Codice prenotazione` — se cambia anche la camera
(`Risorsa`) sono camere realmente distinte della stessa prenotazione multi-camera, da tenere
entrambe; se la camera è la stessa, sono versioni della stessa riga nel tempo, da collassare
tenendo solo quella con `Data cancellazione` più recente (verificato su un caso reale: codice
905, 3 righe, stessa camera, importi e stato diversi — l'ultima per data è quella "buona").
"""

from datetime import date, datetime
from typing import Dict, List

COLONNE_OBBLIGATORIE = {
    "portaleXAlbergo.nome", "dataPren", "arrivo", "partenza", "importo",
}

COLONNE_ELENCO_PRENOTAZIONI = {
    "Codice prenotazione", "Arrivo", "Partenza", "Data prenotazione", "Data cancellazione", "Totale",
}


def _decodifica(raw: bytes) -> str:
    """Prova UTF-8 con BOM, poi UTF-8 puro, poi latin-1 (Windows-1252)."""
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def _hotel_da_portale(nome_portale: str) -> str:
    p = (nome_portale or "").upper()
    if "DU PARC" in p or "DUPARC" in p:
        return "DPH"
    if "CLUB" in p:
        return "CLB"
    if "INTERNATIONAL" in p:
        return "INT"
    return "?"


def _canale_da_portale(nome_portale: str) -> str:
    if " - " in (nome_portale or ""):
        return nome_portale.split(" - ", 1)[1].strip()
    return (nome_portale or "").strip()


def _parse_data(v: str) -> date:
    v = (v or "").strip()
    # Tronca eventuali secondi frazionari (es. "2026-01-17T18:34:56.36")
    v = v.split(".")[0] if "T" in v and "." in v else v
    return datetime.fromisoformat(v).date()


def _num(v: str) -> float:
    v = (v or "").strip()
    return float(v) if v else 0.0


def _num_it(v) -> float:
    """Numero in formato italiano (virgola decimale) — usato solo nel formato
    "Elenco Prenotazioni". Accetta anche un float/int già nativo (celle xlsx numeriche
    non arrivano come stringa)."""
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    v = str(v).strip().replace(".", "").replace(",", ".")
    try:
        return float(v)
    except ValueError:
        return 0.0


def _parse_data_it(v) -> date:
    """Data gg/mm/aaaa (formato "Elenco Prenotazioni"). Accetta anche un oggetto
    date/datetime già nativo (celle xlsx con formattazione data)."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    v = (v or "").strip()
    return datetime.strptime(v, "%d/%m/%Y").date()


def _serializza_grezzo(riga: Dict) -> Dict:
    """Copia JSON-safe della riga originale del file (tutte le colonne, comprese quelle non
    ancora mappate a un campo proprio, es. Segmento/Fonte/Nazione) — le celle xlsx con
    formattazione data/ora arrivano come oggetti `date`/`datetime` nativi, non serializzabili
    direttamente in JSON, vanno convertite a stringa prima di salvarle in `dati_grezzi`."""
    out = {}
    for k, v in riga.items():
        if isinstance(v, (datetime, date)):
            out[k] = v.isoformat()
        elif v is None or isinstance(v, (str, int, float, bool)):
            out[k] = v
        else:
            out[k] = str(v)
    return out


def _hotel_da_ubicazione(ubicazione: str) -> str:
    u = (ubicazione or "").upper()
    if "DU PARC" in u or "DUPARC" in u:
        return "DPH"
    if "CLUB" in u:
        return "CLB"
    if "INTERNATIONAL" in u:
        return "INT"
    return "?"


def _righe_grezze_elenco_prenotazioni_csv(raw: bytes) -> List[Dict]:
    import csv
    import io

    testo = _decodifica(raw)
    reader = csv.DictReader(io.StringIO(testo), delimiter=";")
    return list(reader)


def _righe_grezze_elenco_prenotazioni_xlsx(raw: bytes) -> List[Dict]:
    import io
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    righe_iter = ws.iter_rows(values_only=True)
    intestazioni = [str(c).strip() if c is not None else "" for c in next(righe_iter)]
    risultato = []
    for valori in righe_iter:
        risultato.append({intestazioni[i]: valori[i] for i in range(min(len(intestazioni), len(valori)))})
    return risultato


def _e_formato_elenco_prenotazioni(prima_riga: str) -> bool:
    return "Codice prenotazione" in prima_riga


def parse_elenco_prenotazioni(raw: bytes, righe_grezze: List[Dict], mese_atteso: int, anno_atteso: int, data_rilevata) -> dict:
    """Formato "Elenco Prenotazioni" — filtrato lato Welcome per mese/periodo di ARRIVO (non di
    prenotazione), su un range scelto liberamente dall'utente (anche l'intera stagione, più
    hotel insieme — `Ubicazione` distingue la struttura riga per riga). A differenza del
    formato "PrenotazioniWeb", qui `mese`/`anno` sono solo un'etichetta per l'import (nessun
    controllo di coerenza sulle date): ogni riga porta già arrivo/prenotazione/cancellazione
    reali, non serve validarle contro un singolo mese dichiarato — un "mega import" di più mesi
    non deve generare falsi avvisi.
    Raggruppa per (codice prenotazione, camera) e tiene solo la versione più recente per data
    di cancellazione, per non contare più volte una prenotazione modificata nel tempo (vedi
    docstring in cima al file)."""
    n_totali = 0
    scartate_senza_codice = 0
    gruppi: Dict[tuple, dict] = {}
    warning: List[str] = []

    for riga in righe_grezze:
        n_totali += 1
        codice = str(riga.get("Codice prenotazione") or "").strip()
        if not codice:
            scartate_senza_codice += 1
            continue
        try:
            arrivo = _parse_data_it(riga.get("Arrivo"))
            partenza = _parse_data_it(riga.get("Partenza"))
            data_pren = _parse_data_it(riga.get("Data prenotazione"))
            data_canc = _parse_data_it(riga.get("Data cancellazione"))
        except (ValueError, TypeError):
            warning.append(f"Riga con codice {codice}: data non valida, scartata")
            continue

        risorsa = (str(riga.get("Risorsa") or "")).strip()
        chiave = (codice, risorsa)

        candidato = dict(
            hotel_code=_hotel_da_ubicazione(str(riga.get("Ubicazione") or "")),
            canale=(str(riga.get("Agenzia") or "")).strip() or "altro",
            canale_vendita=(str(riga.get("Canale vendita") or "")).strip(),
            codice_ota=(str(riga.get("Voucher") or "")).strip(),
            numero_prenotazione=codice,
            data_cancellazione=data_canc,
            data_rilevata=data_rilevata,
            data_prenotazione=data_pren,
            arrivo=arrivo,
            partenza=partenza,
            notti=(partenza - arrivo).days,
            pax=int(riga.get("Pax") or 0),
            cliente=(str(riga.get("Ospite") or "")).strip(),
            email=None,
            tipo_camera=risorsa,
            trattamento=(str(riga.get("Trattamento") or "")).strip() or None,
            mercato=(str(riga.get("Mercato") or "")).strip() or None,
            importo=_num_it(riga.get("Totale")),
            dati_grezzi=_serializza_grezzo(riga),
        )

        # Stessa prenotazione+camera vista più volte = versioni storiche: tiene solo quella
        # con data di cancellazione più recente (None trattato come "più vecchia possibile").
        esistente = gruppi.get(chiave)
        if esistente is None or (candidato["data_cancellazione"] or date.min) >= (esistente["data_cancellazione"] or date.min):
            gruppi[chiave] = candidato

    righe = list(gruppi.values())

    if not righe:
        raise ValueError("Nessuna riga valida trovata nel file")

    n_versioni_scartate = n_totali - scartate_senza_codice - len(righe)
    if n_versioni_scartate > 0:
        warning.append(
            f"{n_versioni_scartate} righe erano versioni superate della stessa prenotazione/camera "
            "(stesso codice e stessa camera con dati diversi) — tenuta solo la più recente per data di cancellazione"
        )

    return {
        "righe": righe,
        "n_righe_totali": n_totali,
        "n_righe_valide": len(righe),
        "n_righe_fuori_mese": 0,
        "warning": warning,
    }


def _parse_data_manuale(v: str):
    """Data inserita a mano nella colonna opzionale `data_cancellazione` — accetta sia
    ISO (2026-09-05) sia formato italiano (05/09/2026), per non vincolare a un formato
    scomodo da digitare in Excel. None se assente o non riconosciuta."""
    v = (v or "").strip()
    if not v:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(v, fmt).date()
        except ValueError:
            continue
    return None


def parse_csv(raw: bytes, mese_atteso: int, anno_atteso: int, data_rilevata) -> dict:
    """Punto d'ingresso unico: rileva il formato e smista al parser giusto.

    - Byte iniziali `PK` (firma zip) → file `.xlsx` → sempre formato "Elenco Prenotazioni"
      (l'unico dei due mai visto in xlsx).
    - Altrimenti CSV: se l'intestazione contiene "Codice prenotazione" → "Elenco
      Prenotazioni" (punto e virgola); altrimenti formato "PrenotazioniWeb" (virgola).
    """
    if raw[:2] == b"PK":
        righe_grezze = _righe_grezze_elenco_prenotazioni_xlsx(raw)
        return parse_elenco_prenotazioni(raw, righe_grezze, mese_atteso, anno_atteso, data_rilevata)

    testo_sniff = _decodifica(raw[:2000])
    if _e_formato_elenco_prenotazioni(testo_sniff):
        righe_grezze = _righe_grezze_elenco_prenotazioni_csv(raw)
        return parse_elenco_prenotazioni(raw, righe_grezze, mese_atteso, anno_atteso, data_rilevata)

    return _parse_csv_prenotazioniweb(raw, mese_atteso, anno_atteso, data_rilevata)


def _parse_csv_prenotazioniweb(raw: bytes, mese_atteso: int, anno_atteso: int, data_rilevata) -> dict:
    """Analizza il CSV formato "PrenotazioniWeb" e ritorna dict con `righe`, contatori, `warning`.

    Solleva ValueError solo per errori strutturali (file vuoto, colonne mancanti).
    Righe con `dataPren` fuori dal mese/anno dichiarati vengono incluse comunque
    (Welcome a volte include qualche giorno a cavallo) ma contate a parte, come
    warning — stesso principio del filtro "fuori stagione" del parser Revenue.
    """
    import csv
    import io

    testo = _decodifica(raw)
    righe_testo = testo.splitlines()
    # Scarta l'eventuale riga "sep=," che Excel antepone al CSV
    if righe_testo and righe_testo[0].strip().lower().startswith("sep="):
        righe_testo = righe_testo[1:]
    if not righe_testo:
        raise ValueError("File vuoto")

    reader = csv.DictReader(io.StringIO("\n".join(righe_testo)))
    intestazioni = set(reader.fieldnames or [])
    mancanti = COLONNE_OBBLIGATORIE - intestazioni
    if mancanti:
        raise ValueError(f"Colonne obbligatorie mancanti nel file: {', '.join(sorted(mancanti))}")

    righe: List[Dict] = []
    n_totali = 0
    n_fuori_mese = 0
    warning: List[str] = []

    for riga in reader:
        n_totali += 1
        portale = (riga.get("portaleXAlbergo.nome") or "").strip()
        if not portale:
            continue
        try:
            data_pren = _parse_data(riga.get("dataPren", ""))
            arrivo = _parse_data(riga.get("arrivo", ""))
            partenza = _parse_data(riga.get("partenza", ""))
        except ValueError:
            warning.append(f"Riga {n_totali}: data non valida, scartata")
            continue

        if data_pren.month != mese_atteso or data_pren.year != anno_atteso:
            n_fuori_mese += 1

        righe.append(dict(
            hotel_code=_hotel_da_portale(portale),
            canale=_canale_da_portale(portale) or "altro",
            canale_vendita=(riga.get("parametroCanaleVendita.descrizione") or "").strip(),
            codice_ota=(riga.get("codiceOTA") or "").strip(),
            # Colonne opzionali, non nell'export Welcome originale — compilate solo se l'utente
            # le ha aggiunte a mano al CSV. Accetta sia il nome a schermo in Welcome ("Codice
            # Prenotazione", con spazio) sia lo stile camelCase coerente col resto del file
            # (codicePrenotazione) — verificato che l'utente usa quest'ultimo in pratica.
            numero_prenotazione=(
                riga.get("codicePrenotazione") or riga.get("Codice Prenotazione") or ""
            ).strip() or None,
            data_cancellazione=_parse_data_manuale(
                riga.get("dataCancellazione") or riga.get("Data cancellazione") or ""
            ),
            data_rilevata=data_rilevata,
            data_prenotazione=data_pren,
            arrivo=arrivo,
            partenza=partenza,
            notti=int(riga.get("notti") or 0),
            pax=int(riga.get("pax") or 0),
            cliente=(riga.get("cliente.ragioneSociale") or "").strip(),
            email=(riga.get("cliente.email") or "").strip() or None,
            tipo_camera=(riga.get("tipoCamera.nome") or "").strip(),
            trattamento=(riga.get("trattamento.nome") or "").strip() or None,
            mercato=(riga.get("parametroMercato.descrizione") or "").strip() or None,
            importo=_num(riga.get("importo", "")),
            dati_grezzi=_serializza_grezzo(riga),
        ))

    if not righe:
        raise ValueError("Nessuna riga valida trovata nel file")

    if n_fuori_mese:
        warning.append(
            f"{n_fuori_mese} righe hanno data di prenotazione fuori dal mese/anno dichiarati "
            f"({mese_atteso:02d}/{anno_atteso}) — incluse comunque, verificare il filtro usato in Welcome"
        )

    return {
        "righe": righe,
        "n_righe_totali": n_totali,
        "n_righe_valide": len(righe),
        "n_righe_fuori_mese": n_fuori_mese,
        "warning": warning,
    }
