"""Test integrazione Conto Economico USALI contabile: anteprima/conferma import, reimport,
storico con file originali, report (totali dalle foglie, progressivo, anno precedente,
segmenti) ed export xlsx.

Fixture reali in tests/fixtures/usali/. DB dedicato (TEST_DATABASE_URL), file originali in una
cartella temporanea (USALI_CONTABILE_DIR).
"""
import io
import os
import sys
from types import SimpleNamespace

import openpyxl
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from sqlalchemy import create_engine, text  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402
from starlette.testclient import TestClient  # noqa: E402

from app.auth import richiedi_admin, richiedi_utente_attivo  # noqa: E402
from app.database import Base, get_db  # noqa: E402
from app.main import app  # noqa: E402
from app.models.usali_contabile import UsaliImport, UsaliRiga  # noqa: E402

from tests._db import TEST_DB_URL  # noqa: E402

FIXTURES = os.path.join(os.path.dirname(__file__), "fixtures", "usali")
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

NETTO_2026 = [-6831.55, -7933.76, -22907.78, -44562.45, -24407.67, 7971.45, 67716.00, 129358.28, 9556.32]
RICAVI_2026 = [0, 0, 0, 0, 35317.90, 93546.83, 167606.73, 230782.06, 54530.24]
GOP_2026 = [-6824.80, -7927.66, -22901.03, -44755.91, -25400.92, 5197.65, 67302.26, 128534.70, 11812.85]
EBITDA_2026 = [-6824.80, -7927.66, -22901.03, -44555.91, -24400.92, 7977.98, 67722.75, 129365.03, 9562.85]


def _leggi(nome):
    with open(os.path.join(FIXTURES, nome), "rb") as f:
        return f.read()


XLSX = _leggi("export-1.xlsx")
CSV = _leggi("pivot-1.csv")


def _xlsx_modificato(modifica) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(XLSX))
    modifica(wb.active)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _solo_ago_set_con_settembre_modificato(ws):
    """Export che copre solo agosto–settembre, con +500 € su Camere Individuali a settembre."""
    ws.delete_cols(2, 7)            # via Gennaio..Luglio
    ws.cell(2, 2).value = "2026"    # l'anno stava nella prima colonna mese, cancellata
    ws.cell(10, 3).value = (ws.cell(10, 3).value or 0) + 500   # Settembre, Camere Individuali
    riga_tot = ws.max_row
    ws.cell(riga_tot, 3).value = ws.cell(riga_tot, 3).value + 500


def _anno_2025(ws):
    ws.cell(2, 2).value = "2025"


@pytest.fixture(scope="module")
def test_engine():
    engine = create_engine(TEST_DB_URL)
    Base.metadata.create_all(engine)
    yield engine
    Base.metadata.drop_all(engine)
    engine.dispose()


@pytest.fixture(scope="module")
def TestSession(test_engine):
    return sessionmaker(bind=test_engine)


@pytest.fixture(scope="module")
def client(test_engine, TestSession, tmp_path_factory):
    os.environ["USALI_CONTABILE_DIR"] = str(tmp_path_factory.mktemp("usali_contabile"))

    def override_get_db():
        db = TestSession()
        try:
            yield db
        finally:
            db.close()

    utente_finto = SimpleNamespace(id=None)
    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[richiedi_utente_attivo] = lambda: utente_finto
    app.dependency_overrides[richiedi_admin] = lambda: utente_finto
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()
    os.environ.pop("USALI_CONTABILE_DIR", None)


@pytest.fixture(autouse=True)
def pulisci_e_semina(test_engine):
    with test_engine.begin() as conn:
        conn.execute(text(
            "TRUNCATE usali_riga, usali_import, usali_regola_segmento, usali_segmento, "
            "usali_struttura_gestionale, hotels RESTART IDENTITY CASCADE"
        ))
        conn.execute(text("""
            INSERT INTO hotels (code, name, default_rooms) VALUES
            ('DPH', 'Hotel Du Parc', 43), ('CLB', 'Club Hotel', 45), ('INT', 'International', 45)
        """))
        conn.execute(text(
            "INSERT INTO usali_struttura_gestionale (id_gestionale, struttura_code) VALUES (128, 'DPH')"
        ))
        conn.execute(text("INSERT INTO usali_segmento (struttura_code, nome) VALUES ('DPH', 'Du Parc + Maremosso')"))
        conn.execute(text(
            "INSERT INTO usali_regola_segmento (struttura_code, priorita, tipo_match, pattern, segmento_id, percentuale) "
            "VALUES ('DPH', 0, 'percorso_prefisso', '', 1, 100)"
        ))


def _anteprima(client, xlsx=XLSX, csv=CSV, nome_xlsx="export-1.xlsx", struttura=None):
    files = {}
    if xlsx is not None:
        files["xlsx"] = (nome_xlsx, xlsx, XLSX_MIME)
    if csv is not None:
        files["csv"] = ("pivot-1.csv", csv, "text/csv")
    data = {"struttura_code": struttura} if struttura else {}
    return client.post("/usali/contabile/anteprima", files=files, data=data)


def _importa(client, **kw):
    r = _anteprima(client, **kw)
    assert r.status_code == 200, r.text
    c = client.post("/usali/contabile/conferma", json={"token": r.json()["token"],
                                                       "struttura_code": kw.get("struttura")})
    assert c.status_code == 200, c.text
    return c.json()


def _report(client, **params):
    params = {"struttura": "DPH", "anno": 2026, "mese_da": 1, "mese_a": 9, **params}
    r = client.get("/usali/contabile/report", params=params)
    assert r.status_code == 200, r.text
    return r.json()


# ── Anteprima / conferma ─────────────────────────────────────────────────────

def test_anteprima(client, TestSession):
    r = _anteprima(client)
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["struttura_rilevata"] == "DPH" and j["struttura"] == "DPH"
    assert j["quadratura_ok"] and j["csv_arricchito"]
    assert j["n_conti"] == 223
    assert j["periodi"] == [f"2026-{m:02d}" for m in range(1, 10)]
    assert all(not c["gia_presente"] for c in j["confronto"])
    assert [c["netto_nuovo"] for c in j["confronto"]] == NETTO_2026
    assert j["gia_importato"] is None
    with TestSession() as db:   # l'anteprima non scrive nulla
        assert db.query(UsaliRiga).count() == 0


def test_conferma_e_report(client, TestSession):
    res = _importa(client)
    assert res["righe_inserite"] == 223 * 9 and res["righe_sostituite"] == 0
    with TestSession() as db:
        assert db.query(UsaliRiga).count() == 2007
        imp = db.query(UsaliImport).one()
        assert os.path.exists(imp.percorso_xlsx) and os.path.exists(imp.percorso_csv)

    rep = _report(client)
    s = rep["sintesi"]
    assert s["netto"]["valori"] == NETTO_2026
    assert s["netto"]["totale"] == 107958.84
    assert s["ricavi"]["valori"] == RICAVI_2026
    assert s["ricavi"]["totale"] == 581783.76
    assert s["gop"]["valori"] == GOP_2026
    assert s["gop"]["totale"] == 105037.14
    assert s["ebitda"]["valori"] == EBITDA_2026
    assert s["ricavi"]["perc_ricavi"] == 100.0
    assert rep["prec_disponibile"] is False and s["netto"]["anno_prec"] is None


def test_profitto_per_reparto(client):
    _importa(client)
    righe = {r["id"]: r for r in _report(client)["righe"]}
    camere = righe["profitto:Camere"]
    ric = righe["rep:Ricavi Operativi|Camere"]["totale"]
    cos = righe["rep:Costi Operativi|Camere"]["totale"]
    assert camere["totale"] == pytest.approx(ric + cos, abs=0.01)
    # Reparto solo-ricavi: il profitto coincide con i ricavi
    assert righe["profitto:Ricavi Vari"]["totale"] == righe["rep:Ricavi Operativi|Ricavi Vari"]["totale"]
    # Somma dei profitti per reparto = Profitto dei Reparti Operativi
    somma = sum(r["totale"] for r in righe.values() if r["tipo"] == "profitto_reparto")
    assert somma == pytest.approx(righe["tot:profitto_reparti"]["totale"], abs=0.05)
    # Stesso nome conto in due punti → due righe distinte
    altri = [r for r in righe.values() if r["tipo"] == "conto" and r["nome"] == "Altri Alimenti"]
    assert len(altri) == 6


def test_mesi_non_importati_vuoti_e_progressivo(client):
    _importa(client)
    rep = _report(client, mese_da=6, mese_a=12)
    netto = rep["sintesi"]["netto"]
    assert netto["valori"][:4] == NETTO_2026[5:9]
    assert netto["valori"][4:] == [None, None, None]   # ott–dic: non importati, non zero
    assert rep["mesi_importati"] == [6, 7, 8, 9]
    assert netto["progressivo"] == 107958.84           # gen–dic, da inizio anno
    assert netto["totale"] == pytest.approx(sum(NETTO_2026[5:9]), abs=0.01)


def test_stesso_file_reimportato(client, TestSession):
    _importa(client)
    j = _anteprima(client).json()
    assert j["gia_importato"] is not None
    assert all(c["differenza"] == 0 for c in j["confronto"])
    c = client.post("/usali/contabile/conferma", json={"token": j["token"]})
    assert c.status_code == 409
    with TestSession() as db:
        assert db.query(UsaliImport).count() == 1
        assert db.query(UsaliRiga).count() == 2007


def test_reimport_mese_modificato(client, TestSession):
    primo = _importa(client)
    nuovo = _xlsx_modificato(_solo_ago_set_con_settembre_modificato)

    j = _anteprima(client, xlsx=nuovo, csv=None, nome_xlsx="export-2.xlsx", struttura="DPH").json()
    assert j["periodi"] == ["2026-08", "2026-09"]
    conf = {c["periodo"]: c for c in j["confronto"]}
    assert conf["2026-08"]["differenza"] == 0
    assert conf["2026-09"]["differenza"] == 500
    assert conf["2026-09"]["netto_vecchio"] == 9556.32

    c = client.post("/usali/contabile/conferma", json={"token": j["token"], "struttura_code": "DPH"})
    assert c.status_code == 200, c.text
    assert c.json()["righe_sostituite"] == 223 * 2

    with TestSession() as db:
        assert db.query(UsaliRiga).count() == 2007
        per_import = dict(
            db.execute(text("SELECT import_id, count(DISTINCT mese) FROM usali_riga GROUP BY import_id")).all()
        )
        assert per_import == {primo["import_id"]: 7, c.json()["import_id"]: 2}

    netto = _report(client)["sintesi"]["netto"]["valori"]
    assert netto[:8] == NETTO_2026[:8]
    assert netto[8] == pytest.approx(NETTO_2026[8] + 500, abs=0.01)

    storico = client.get("/usali/contabile/import").json()
    assert len(storico) == 2   # il log conserva anche il primo import
    vecchio = next(i for i in storico if i["id"] == primo["import_id"])
    assert vecchio["periodi_in_vigore"] == [f"2026-{m:02d}" for m in range(1, 8)]
    d = client.get(f"/usali/contabile/import/{primo['import_id']}/file/xlsx")
    assert d.status_code == 200 and d.content == XLSX


def test_solo_csv_rifiutato(client):
    r = _anteprima(client, xlsx=None)
    assert r.status_code == 400 and "xlsx" in r.json()["detail"]
    r = _anteprima(client, xlsx=CSV, nome_xlsx="pivot-1.csv")
    assert r.status_code == 400 and "xlsx" in r.json()["detail"]


def test_senza_csv_struttura_da_scegliere(client):
    j = _anteprima(client, csv=None).json()
    assert j["struttura_rilevata"] is None and j["struttura"] is None
    assert not j["csv_arricchito"]
    c = client.post("/usali/contabile/conferma", json={"token": j["token"]})
    assert c.status_code == 400
    c = client.post("/usali/contabile/conferma", json={"token": j["token"], "struttura_code": "DPH"})
    assert c.status_code == 200


def test_struttura_in_conflitto_col_csv(client):
    r = _anteprima(client, struttura="CLB")
    assert r.status_code == 400 and "DPH" in r.json()["detail"]


def test_confronto_anno_precedente(client):
    _importa(client)
    _importa(client, xlsx=_xlsx_modificato(_anno_2025), csv=None, struttura="DPH")
    rep = _report(client)
    netto = rep["sintesi"]["netto"]
    assert rep["prec_disponibile"]
    assert netto["anno_prec"] == 107958.84
    assert netto["delta_anno_prec"] == 0
    # Il 2025 non entra nel progressivo 2026
    assert netto["progressivo"] == 107958.84


def test_segmento(client, test_engine):
    _importa(client)
    totale = _report(client)["sintesi"]
    seg = _report(client, segmento_id=1)
    assert seg["errori_segmento"] == []
    assert seg["sintesi"]["netto"]["valori"] == totale["netto"]["valori"]

    with test_engine.begin() as conn:
        conn.execute(text("INSERT INTO usali_segmento (struttura_code, nome) VALUES ('DPH', 'Maremosso')"))
        conn.execute(text("""
            INSERT INTO usali_regola_segmento (struttura_code, priorita, tipo_match, pattern, segmento_id, percentuale) VALUES
            ('DPH', 10, 'percorso_prefisso', 'Food & Beverage > Ristorante', 2, 100),
            ('DPH', 10, 'nome_contiene', 'Maremosso', 2, 100),
            ('DPH', 10, 'percorso_prefisso', 'Costi Generali > Utenze', 1, 70),
            ('DPH', 10, 'percorso_prefisso', 'Costi Generali > Utenze', 2, 30)
        """))
    hotel = _report(client, segmento_id=1)["sintesi"]
    mms = _report(client, segmento_id=2)["sintesi"]
    for k in ("ricavi", "gop", "ebitda", "netto"):
        for h, m, t in zip(hotel[k]["valori"], mms[k]["valori"], totale[k]["valori"]):
            # ogni segmento è arrotondato al centesimo a sé: ±1 cent (esattezza a 4 decimali
            # già verificata in test_usali_contabile_parser.py)
            assert h + m == pytest.approx(t, abs=0.015)

    # Regola che lascia scoperte delle righe → errore visibile nel report
    with test_engine.begin() as conn:
        conn.execute(text("DELETE FROM usali_regola_segmento WHERE priorita = 0"))
    rep = _report(client, segmento_id=1)
    assert rep["errori_segmento"]
    assert all("nessuna regola" in e["motivo"] for e in rep["errori_segmento"])


def test_segmento_di_altra_struttura(client):
    _importa(client)
    r = client.get("/usali/contabile/report", params={"struttura": "CLB", "anno": 2026, "segmento_id": 1})
    assert r.status_code == 404


def test_export_xlsx(client):
    _importa(client)
    r = client.get("/usali/contabile/report/export",
                   params={"struttura": "DPH", "anno": 2026, "mese_da": 1, "mese_a": 9})
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    voci = [ws.cell(i, 1).value for i in range(1, ws.max_row + 1)]
    riga_netto = voci.index("Reddito Netto") + 1
    assert ws.cell(riga_netto, 11).value == pytest.approx(107958.84, abs=0.01)   # Totale periodo


def test_strutture_e_periodi(client):
    _importa(client)
    strutture = {s["code"]: s for s in client.get("/usali/contabile/strutture").json()}
    assert strutture["DPH"]["ha_dati"] and strutture["DPH"]["id_gestionale"] == [128]
    assert "BON" in strutture and not strutture["CLB"]["ha_dati"]
    periodi = client.get("/usali/contabile/periodi", params={"struttura": "DPH"}).json()
    assert periodi == [{"anno": 2026, "mesi": list(range(1, 10))}]
