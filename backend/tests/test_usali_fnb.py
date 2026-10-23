"""Test USALI → Food & Beverage.

Calcolo puro (`costruisci_report`) su dati sintetici + endpoint coperti/parametri sul DB di test.
"""
import os
import sys
from types import SimpleNamespace

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from sqlalchemy import create_engine, text  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402
from starlette.testclient import TestClient  # noqa: E402

from app.auth import richiedi_admin, richiedi_utente_attivo  # noqa: E402
from app.database import Base, get_db  # noqa: E402
from app.main import app  # noqa: E402
from app.services import usali_fnb as fnb  # noqa: E402

from tests._db import TEST_DB_URL  # noqa: E402


def _mese(importato=True, contabile=None, personale=None, ricavi_prod=None, pasti=None,
          presenze=None, esterni=0.0, bon=0.0, menu=None, scontrini=None):
    return {
        'importato': importato, 'presenze': presenze,
        'contabile': contabile or {}, 'personale': personale or {},
        'ricavi_prod': ricavi_prod or {}, 'pasti': pasti or {},
        'ricavi_esterni': esterni, 'ricavi_bon': bon,
        'coperti_menu': menu, 'scontrini_bon': scontrini,
    }


def _cont(food=0.0, beverage=0.0, altri=0.0, ricavi=0.0, personale_contabile=0.0):
    return {'ricavi': 0.0, 'ricavi_contabili': ricavi, 'food': food, 'beverage': beverage,
            'personale': 0.0, 'personale_contabile': personale_contabile, 'altri': altri}


def _dati(struttura, mesi):
    tutti = {m: _mese(importato=False) for m in range(1, 13)}
    tutti.update(mesi)
    return {'struttura': struttura, 'anno': 2026, 'mesi': tutti, 'conti': {}}


def _riga(rep, rid):
    return next(r for r in rep['righe'] if r['id'] == rid)


# ── Calcolo puro ─────────────────────────────────────────────────────────────

def test_ristorante_hotel_riga_unica_pranzi_cene_informativi():
    dati = _dati('CLB', {6: _mese(
        contabile={'ristorante': _cont(food=900.0, beverage=100.0, altri=200.0, ricavi=4000.0)},
        personale={'ristorante': 1000.0},
        ricavi_prod={('ristorante', 'pranzo'): 1000.0, ('ristorante', 'cena'): 3000.0,
                     ('ristorante', None): 100.0},
        pasti={'pranzo': 100, 'cena': 100}, presenze=400,
    )})
    rep = fnb.costruisci_report(dati, 6, 6)
    assert [r['id'] for r in rep['righe'] if r['punto'] == 'ristorante'] == ['CLB.ristorante']
    rist = _riga(rep, 'CLB.ristorante')['periodo']
    assert rist['ricavi'] == 4100.0          # pranzo + cena + voci non smistate
    assert rist['materia_prima'] == 1000.0
    assert rist['coperti'] == 200
    assert rist['costo_coperto'] == 11.0     # (1000 + 1000 + 200) / 200, uguale per pranzo e cena
    assert (rist['food_coperto'], rist['beverage_coperto']) == (4.5, 0.5)
    assert (rist['pranzi'], rist['cene']) == (100, 100)
    assert (rist['pranzi_perc'], rist['cene_perc']) == (25.0, 25.0)


def test_kpi_di_periodo_sui_totali_non_media_dei_mesi():
    dati = _dati('CLB', {
        6: _mese(contabile={'colazione': _cont(food=100.0)}, ricavi_prod={('colazione', None): 1000.0},
                 pasti={'colazione': 100}),
        7: _mese(contabile={'colazione': _cont(food=900.0)}, ricavi_prod={('colazione', None): 3000.0},
                 pasti={'colazione': 300}),
    })
    p = _riga(fnb.costruisci_report(dati, 6, 7), 'CLB.colazione')['periodo']
    assert p['food_perc'] == 25.0            # 1000/4000, non media di 10% e 30%
    assert p['mp_coperto'] == 2.5            # 1000/400


def test_mesi_senza_contabilita_esclusi_dai_rapporti():
    dati = _dati('CLB', {
        6: _mese(contabile={'colazione': _cont(food=300.0)}, ricavi_prod={('colazione', None): 1000.0},
                 pasti={'colazione': 100}),
        7: _mese(importato=False, ricavi_prod={('colazione', None): 5000.0}, pasti={'colazione': 500}),
    })
    rep = fnb.costruisci_report(dati, 6, 7)
    p = _riga(rep, 'CLB.colazione')['periodo']
    assert p['ricavi'] == 6000.0             # i ricavi contano sempre
    assert p['food_perc'] == 30.0            # solo giugno, unico mese con costi
    assert p['mp_coperto'] == 3.0
    assert p['ricavo_coperto'] == 10.0       # ricavi e coperti su entrambi i mesi
    assert rep['mesi_senza_contabilita'] == [7]
    luglio = _riga(rep, 'CLB.colazione')['mesi'][1]
    assert luglio['food'] is None and luglio['food_perc'] is None


def test_costo_per_pasto_scomposto_e_margine():
    dati = _dati('CLB', {
        6: _mese(contabile={'colazione': _cont(food=200.0, altri=50.0, ricavi=900.0)},
                 personale={'colazione': 250.0},
                 ricavi_prod={('colazione', None): 1000.0}, pasti={'colazione': 100}),
        7: _mese(importato=False, ricavi_prod={('colazione', None): 3000.0}, pasti={'colazione': 300}),
    })
    p = _riga(fnb.costruisci_report(dati, 6, 7), 'CLB.colazione')['periodo']
    assert (p['mp_coperto'], p['personale_coperto'], p['altri_coperto']) == (2.0, 2.5, 0.5)
    assert p['costo_coperto'] == 5.0
    assert p['ricavo_coperto'] == 10.0          # su tutti i mesi con coperti
    assert p['ricavo_coperto_cc'] == 10.0       # sui soli mesi con costi
    assert p['margine_coperto'] == 5.0


def test_costi_dei_mesi_senza_vendite_nel_costo_per_pasto():
    dati = _dati('CLB', {
        4: _mese(contabile={'colazione': _cont(food=100.0)}),     # acquisti prima dell'apertura
        6: _mese(contabile={'colazione': _cont(food=300.0)}, ricavi_prod={('colazione', None): 1000.0},
                 pasti={'colazione': 100}),
        7: _mese(contabile={'colazione': _cont(food=900.0)}, ricavi_prod={('colazione', None): 500.0}),
    })
    p = _riga(fnb.costruisci_report(dati, 4, 7), 'CLB.colazione')['periodo']
    # aprile (nessuna vendita) entra, luglio (vendite senza pasti = dato mancante) no
    assert p['mp_coperto'] == 4.0
    assert p['margine_coperto'] == 6.0


def test_mare_mosso_costi_avvio_ripartiti_sulle_quote_di_periodo():
    dati = _dati('DPH', {
        4: _mese(contabile={'ristorante': _cont(food=400.0)}),
        8: _mese(contabile={'ristorante': _cont(food=1200.0, ricavi=5000.0)},
                 ricavi_prod={('ristorante', 'ospiti'): 2000.0}, pasti={'cena': 100},
                 esterni=3000.0, menu=400),
    })
    rep = fnb.costruisci_report(dati, 4, 8)
    rist = _riga(rep, 'DPH.ristorante')['periodo']
    ospiti = _riga(rep, 'DPH.ristorante.ospiti')['periodo']
    esterni = _riga(rep, 'DPH.ristorante.esterni')['periodo']
    assert rist['materia_prima'] == 1600.0 and rist['mp_coperto'] == 4.0
    assert round(ospiti['materia_prima'] + esterni['materia_prima'], 2) == 1600.0
    assert ospiti['mp_coperto'] == 4.0 and esterni['mp_coperto'] == 4.0


def test_mare_mosso_esterni_da_coperti_menu():
    dati = _dati('DPH', {8: _mese(
        contabile={'ristorante': _cont(food=1000.0, ricavi=5000.0)},
        ricavi_prod={('ristorante', 'ospiti'): 2000.0}, pasti={'pranzo': 20, 'cena': 80},
        esterni=3000.0, menu=300,
    )})
    rep = fnb.costruisci_report(dati, 8, 8)
    assert _riga(rep, 'DPH.ristorante')['periodo']['coperti'] == 300
    ospiti = _riga(rep, 'DPH.ristorante.ospiti')['periodo']
    esterni = _riga(rep, 'DPH.ristorante.esterni')['periodo']
    assert ospiti['coperti'] == 100 and esterni['coperti'] == 200
    assert esterni['ricavi'] == 3000.0
    assert round(ospiti['materia_prima'] + esterni['materia_prima'], 2) == 1000.0
    assert round(esterni['materia_prima'], 2) == 666.67   # 200 su 300 coperti
    assert rep['avvisi'] == []


def test_mare_mosso_senza_coperti_menu_nessuna_stima():
    dati = _dati('DPH', {8: _mese(
        contabile={'ristorante': _cont(food=1000.0)},
        ricavi_prod={('ristorante', 'ospiti'): 2000.0}, pasti={'cena': 80}, esterni=3000.0,
    )})
    rep = fnb.costruisci_report(dati, 8, 8)
    assert _riga(rep, 'DPH.ristorante')['periodo']['materia_prima'] == 1000.0
    assert _riga(rep, 'DPH.ristorante.esterni')['periodo']['materia_prima'] is None
    assert any('coperti Menu non inseriti' in a for a in rep['avvisi'])


def test_mare_mosso_menu_inferiore_agli_ospiti_segnalato():
    dati = _dati('DPH', {8: _mese(ricavi_prod={('ristorante', 'ospiti'): 2000.0},
                                  pasti={'cena': 80}, esterni=100.0, menu=50)})
    rep = fnb.costruisci_report(dati, 8, 8)
    assert _riga(rep, 'DPH.ristorante.esterni')['periodo']['coperti'] is None
    assert any('inferiori ai pasti' in a for a in rep['avvisi'])


def test_bon_scontrini_e_avviso_contabilita_incompleta():
    dati = _dati('BON', {
        7: _mese(contabile={'unico': _cont(food=500.0, ricavi=0.0)}, bon=10000.0, scontrini=1000),
    })
    rep = fnb.costruisci_report(dati, 7, 7)
    p = _riga(rep, 'BON.unico')['periodo']
    assert p['ricavo_coperto'] == 10.0 and p['mp_coperto'] == 0.5
    assert any('contabilità probabilmente incompleta' in a for a in rep['avvisi'])


def test_totale_struttura_e_gruppo_sommano_le_basi():
    a = fnb.costruisci_report(_dati('CLB', {6: _mese(
        contabile={'colazione': _cont(food=100.0), 'bar': _cont(beverage=50.0)},
        ricavi_prod={('colazione', None): 1000.0, ('bar', None): 500.0})}), 6, 6)
    b = fnb.costruisci_report(_dati('INT', {6: _mese(
        contabile={'colazione': _cont(food=300.0)}, ricavi_prod={('colazione', None): 500.0})}), 6, 6)
    assert a['righe'][0]['periodo']['materia_prima'] == 150.0
    assert a['righe'][0]['periodo']['coperti'] is None
    g = fnb.report_gruppo([a, b], 6, 6)
    tot = g['righe'][0]['periodo']
    assert tot['ricavi'] == 2000.0 and tot['materia_prima'] == 450.0
    assert tot['materia_prima_perc'] == 22.5


def test_classificazioni():
    assert fnb._punto_produzione('CLB', 'Ristorante', 'cena') == ('ristorante', 'cena')
    assert fnb._punto_produzione('DPH', 'Ristorante', 'pranzo') == ('ristorante', 'ospiti')
    assert fnb._punto_produzione('DPH', 'Maremosso', 'altro') == ('ristorante', 'ospiti')
    assert fnb._punto_produzione('CLB', 'Maremosso', 'altro') is None
    assert fnb._punto_produzione('INT', 'Hotel', 'colazione_extra') == ('colazione', None)
    assert fnb._punto_produzione('INT', 'Bar', 'bar_caffetteria') == ('bar', None)
    assert fnb._punto_produzione('INT', 'Hotel', 'alloggio') is None
    assert fnb.punto_contabile('DPH', 'Ristorante > Ristorante Mare Mosso') == 'ristorante'
    assert fnb.punto_contabile('BON', 'Bar') == 'unico'
    assert fnb.punto_default_cc('CLB_PASTICCERIA') == 'colazione'
    assert fnb.punto_default_cc('INT_BAR') == 'bar'
    assert fnb.punto_default_cc('MMS_BAR') == 'ristorante'
    assert fnb.punto_default_cc('CLB_SALA') == 'ristorante'


def test_valida_parametri():
    assert fnb.valida_parametri({'cc_punto': {'CLB_BAR': 'ristorante'}}) == {'cc_punto': {'CLB_BAR': 'ristorante'}}
    with pytest.raises(ValueError):
        fnb.valida_parametri({'cc_punto': {'CLB_BAR': 'cucina'}})


# ── Endpoint ─────────────────────────────────────────────────────────────────

@pytest.fixture(scope="module")
def client():
    engine = create_engine(TEST_DB_URL)
    Base.metadata.create_all(engine)
    TestSession = sessionmaker(bind=engine)
    with engine.begin() as conn:
        conn.execute(text("TRUNCATE usali_fnb_coperti"))
        conn.execute(text("DELETE FROM app_config WHERE key = 'usali_fnb_parametri'"))

    def override_get_db():
        db = TestSession()
        try:
            yield db
        finally:
            db.close()

    utente = SimpleNamespace(id=None)
    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[richiedi_utente_attivo] = lambda: utente
    app.dependency_overrides[richiedi_admin] = lambda: utente
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()
    Base.metadata.drop_all(engine)
    engine.dispose()


def test_coperti_upsert_e_cancellazione(client):
    r = client.put("/usali/fnb/coperti", json={'struttura_code': 'MMS', 'anno': 1901, 'mese': 8, 'valore': 1200})
    assert r.status_code == 200
    client.put("/usali/fnb/coperti", json={'struttura_code': 'MMS', 'anno': 1901, 'mese': 8, 'valore': 1300})
    client.put("/usali/fnb/coperti", json={'struttura_code': 'BON', 'anno': 1901, 'mese': 8, 'valore': 900})
    mesi = client.get("/usali/fnb/coperti", params={'anno': 1901}).json()['mesi']
    agosto = mesi[7]
    assert agosto['mms_coperti'] == 1300 and agosto['bon_scontrini'] == 900
    assert agosto['ospiti_dph'] == 0 and agosto['mms_esterni'] == 1300
    client.put("/usali/fnb/coperti", json={'struttura_code': 'MMS', 'anno': 1901, 'mese': 8, 'valore': None})
    assert client.get("/usali/fnb/coperti", params={'anno': 1901}).json()['mesi'][7]['mms_coperti'] is None


@pytest.mark.parametrize("body", [
    {'struttura_code': 'DPH', 'anno': 1901, 'mese': 8, 'valore': 1},
    {'struttura_code': 'MMS', 'anno': 1901, 'mese': 13, 'valore': 1},
    {'struttura_code': 'MMS', 'anno': 1901, 'mese': 8, 'valore': -5},
    {'struttura_code': 'MMS', 'anno': 1901, 'mese': 8, 'valore': 'tanti'},
])
def test_coperti_valori_non_validi(client, body):
    assert client.put("/usali/fnb/coperti", json=body).status_code == 400


def test_parametri_salvataggio(client):
    assert 'centri' in client.get("/usali/fnb/parametri").json()
    r = client.put("/usali/fnb/parametri", json={'cc_punto': {'CLB_PASTICCERIA': 'ristorante'}})
    assert r.status_code == 200
    assert client.put("/usali/fnb/parametri", json={'cc_punto': {'CLB_BAR': 'cucina'}}).status_code == 400


def test_report_struttura_non_valida(client):
    assert client.get("/usali/fnb/report", params={'struttura': 'XXX', 'anno': 2026}).status_code == 400
    r = client.get("/usali/fnb/report", params={'struttura': 'GRUPPO', 'anno': 1901})
    assert r.status_code == 200
    assert r.json()['righe'][0]['id'] == 'G.totale'


def test_export_xlsx(client):
    import io
    import openpyxl
    r = client.get("/usali/fnb/report/export", params={'struttura': 'CLB', 'anno': 1901, 'mese_da': 5, 'mese_a': 9})
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert ws.cell(3, 1).value == 'Punto vendita'
    assert ws.cell(4, 1).value == 'Totale Club Hotel'


def test_costo_pasto_endpoint(client):
    r = client.get("/usali/fnb/costo-pasto", params={'anno': 1901, 'mese_da': 5, 'mese_a': 9})
    assert r.status_code == 200
    ids = [x['id'] for x in r.json()['righe']]
    assert ids[:3] == ['DPH.colazione', 'CLB.colazione', 'INT.colazione']
    assert 'CLB.ristorante' in ids and 'DPH.ristorante.esterni' in ids
    assert not any(i.endswith(('.pranzo', '.cena')) for i in ids)
