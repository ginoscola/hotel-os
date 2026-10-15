"""Test del parser export gestionale (Conto Economico USALI contabile) e dei segmenti.

Fixture reali: tests/fixtures/usali/export-1.xlsx + pivot-1.csv (Hotel Du Parc, gen–set 2026).
Test puri, nessun database.
"""
import io
import os
import sys
from decimal import Decimal

import openpyxl
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app.services.usali_contabile_parser import ErroreImportUsali, parse_export  # noqa: E402
from app.services.usali_segmenti import assegna_segmenti  # noqa: E402

FIXTURES = os.path.join(os.path.dirname(__file__), "fixtures", "usali")
PERIODI_2026 = [(2026, m) for m in range(1, 10)]


@pytest.fixture(scope="module")
def xlsx_bytes():
    with open(os.path.join(FIXTURES, "export-1.xlsx"), "rb") as f:
        return f.read()


@pytest.fixture(scope="module")
def csv_bytes():
    with open(os.path.join(FIXTURES, "pivot-1.csv"), "rb") as f:
        return f.read()


@pytest.fixture(scope="module")
def risultato(xlsx_bytes, csv_bytes):
    return parse_export(xlsx_bytes, csv_bytes)


def _totali(righe, filtro=lambda r: True):
    return [
        sum((r.importi[p] for r in righe if filtro(r)), Decimal(0)).quantize(Decimal("0.01"))
        for p in PERIODI_2026
    ]


def _d(*valori):
    return [Decimal(v) for v in valori]


SEZ_GOP = ("Ricavi Operativi", "Costi Operativi", "Costi Generali")


# ── Parsing xlsx ─────────────────────────────────────────────────────────────

def test_foglie_e_periodi(risultato):
    assert len(risultato.righe) == 223
    assert risultato.periodi == PERIODI_2026
    # Righe a zero salvate per ogni mese presente nel file
    assert all(set(r.importi) == set(PERIODI_2026) for r in risultato.righe)


def test_stesso_conto_percorsi_diversi(risultato):
    altri = {r.percorso_testo: r for r in risultato.righe if r.conto_nome == "Altri Alimenti"}
    colazione = [r for p, r in altri.items() if "Ricavi Operativi > Food & Beverage > Colazione" in p]
    ristorante = [r for p, r in altri.items()
                  if "Ricavi Operativi > Food & Beverage > Ristorante > Ristorante Mare Mosso" in p]
    assert len(colazione) == 1 and len(ristorante) == 1
    assert sum(colazione[0].importi.values()) != sum(ristorante[0].importi.values())
    assert ristorante[0].reparto == "Food & Beverage"
    assert ristorante[0].sottoreparto == "Ristorante > Ristorante Mare Mosso"
    assert colazione[0].sottoreparto == "Colazione"


def test_sezioni(risultato):
    sezioni = {r.sezione for r in risultato.righe}
    assert sezioni == {"Ricavi Operativi", "Costi Operativi", "Costi Generali", "Non Operativi", "Deprezzamenti"}
    deprezz = [r for r in risultato.righe if r.sezione == "Deprezzamenti"]
    assert len(deprezz) == 2 and all(r.ammortamento for r in deprezz)
    # Una rettifica di costo positiva resta positiva (nessuna "correzione" di segno)
    non_op = next(r for r in risultato.righe if r.sezione == "Non Operativi")
    assert sum(non_op.importi.values()) > 0


def test_quadratura_ok(risultato):
    assert risultato.quadratura_ok
    assert risultato.differenze_quadratura == {}


def test_valori_attesi(risultato):
    righe = risultato.righe
    assert _totali(righe) == _d("-6831.55", "-7933.76", "-22907.78", "-44562.45", "-24407.67",
                                "7971.45", "67716.00", "129358.28", "9556.32")
    assert sum(_totali(righe)) == Decimal("107958.84")

    ricavi = _totali(righe, lambda r: r.sezione == "Ricavi Operativi")
    assert ricavi == _d("0", "0", "0", "0", "35317.90", "93546.83", "167606.73", "230782.06", "54530.24")
    assert sum(ricavi) == Decimal("581783.76")

    gop = _totali(righe, lambda r: r.sezione in SEZ_GOP)
    assert gop == _d("-6824.80", "-7927.66", "-22901.03", "-44755.91", "-25400.92",
                     "5197.65", "67302.26", "128534.70", "11812.85")
    assert sum(gop) == Decimal("105037.14")

    ebitda = _totali(righe, lambda r: r.sezione != "Deprezzamenti")
    assert ebitda == _d("-6824.80", "-7927.66", "-22901.03", "-44555.91", "-24400.92",
                        "7977.98", "67722.75", "129365.03", "9562.85")


def _xlsx_modificato(xlsx_bytes, modifica):
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    modifica(wb.active)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_quadratura_ko_rifiuta(xlsx_bytes):
    # Camere Individuali, colonna Settembre: +100 € non riflessi nel Gran Totale
    def modifica(ws):
        ws.cell(10, 10).value = (ws.cell(10, 10).value or 0) + 100
    with pytest.raises(ErroreImportUsali, match="non quadra"):
        parse_export(_xlsx_modificato(xlsx_bytes, modifica))


def test_solo_csv_rifiutato(csv_bytes):
    with pytest.raises(ErroreImportUsali, match="xlsx"):
        parse_export(csv_bytes)
    with pytest.raises(ErroreImportUsali, match="xlsx"):
        parse_export(b"", csv_bytes)


# ── Arricchimento CSV ────────────────────────────────────────────────────────

def test_arricchimento_csv(risultato):
    assert risultato.csv_arricchito
    assert risultato.id_struttura_gestionale == 128
    assert risultato.avvisi == []
    assert all(r.id_conti is not None for r in risultato.righe)
    chiavi = [(r.id_conti, r.id_reparto, r.ammortamento) for r in risultato.righe]
    assert len(set(chiavi)) == len(chiavi)
    # (id_conti, id_reparto) da sola non è univoca
    assert len({(r.id_conti, r.id_reparto) for r in risultato.righe}) < len(chiavi)
    costo = next(r for r in risultato.righe
                 if r.sezione == "Costi Operativi" and r.conto_nome == "Acquisto e Noleggio Attrezzature")
    assert costo.tipo == "costo" and costo.id_conti == 4176 and costo.id_reparto == 369
    assert all(r.tipo == "ricavo" for r in risultato.righe if r.sezione == "Ricavi Operativi")


def test_csv_nome_non_coincidente(xlsx_bytes, csv_bytes):
    testo = csv_bytes.decode("utf-8-sig").replace("Camere Individuali|", "Camere Singole|", 1)
    r = parse_export(xlsx_bytes, testo.encode("utf-8-sig"))
    assert not r.csv_arricchito
    assert r.id_struttura_gestionale is None
    assert any("non corrispondono" in a for a in r.avvisi)
    assert all(riga.id_conti is None for riga in r.righe)
    assert len(r.righe) == 223   # l'import procede con i soli dati xlsx


# ── Segmenti ─────────────────────────────────────────────────────────────────

HOTEL, MMS = 1, 2


def _righe_db(risultato):
    """Come le leggerebbe il report da usali_riga: una per conto × mese."""
    return [
        {"struttura_code": "DPH", "anno": a, "mese": m, "percorso": r.percorso,
         "conto_nome": r.conto_nome, "id_conti": r.id_conti, "id_reparto": r.id_reparto,
         "importo": r.importi[(a, m)]}
        for r in risultato.righe for (a, m) in risultato.periodi
    ]


def _regola(tipo, pattern, seg, pct=100, prio=10):
    return {"struttura_code": "DPH", "priorita": prio, "tipo_match": tipo, "pattern": pattern,
            "segmento_id": seg, "percentuale": pct, "valida_dal": "2000-01", "valida_al": None}


def _per_mese(quote, seg=None):
    out = {}
    for q in quote:
        if seg is None or q.segmento_id == seg:
            k = (q.riga["anno"], q.riga["mese"])
            out[k] = out.get(k, Decimal(0)) + q.importo
    return out


def _per_mese_righe(righe):
    out = {}
    for r in righe:
        k = (r["anno"], r["mese"])
        out[k] = out.get(k, Decimal(0)) + r["importo"]
    return out


def test_segmento_catch_all(risultato):
    righe = _righe_db(risultato)
    quote, errori = assegna_segmenti(righe, [_regola("percorso_prefisso", "", HOTEL, prio=0)])
    assert errori == []
    assert _per_mese(quote, HOTEL) == _per_mese_righe(righe)


REGOLE_PROVA = [
    _regola("percorso_prefisso", "", HOTEL, prio=0),
    _regola("percorso_prefisso", "Food & Beverage > Ristorante", MMS),
    _regola("nome_contiene", "maremosso", MMS),
    _regola("percorso_prefisso", "Costi Generali > Utenze", HOTEL, pct=70),
    _regola("percorso_prefisso", "Costi Generali > Utenze", MMS, pct=30),
]


def test_segmenti_hotel_maremosso(risultato):
    righe = _righe_db(risultato)
    quote, errori = assegna_segmenti(righe, REGOLE_PROVA)
    assert errori == []
    hotel, mms, totale = _per_mese(quote, HOTEL), _per_mese(quote, MMS), _per_mese_righe(righe)
    for periodo in risultato.periodi:
        assert hotel[periodo] + mms[periodo] == totale[periodo]

    # Utenze ripartite 70/30
    elettr = [q for q in quote if q.riga["conto_nome"] == "Elettricità" and q.riga["mese"] == 8]
    assert len(elettr) == 2
    imp = next(r["importo"] for r in righe if r["conto_nome"] == "Elettricità" and r["mese"] == 8)
    q_mms = next(q.importo for q in elettr if q.segmento_id == MMS)
    assert q_mms == (imp * Decimal("0.3")).quantize(Decimal("0.0001"))

    # Ricavi e costi del Ristorante + conti "…Maremosso" in Costi Generali vanno a Maremosso
    for q in quote:
        nome, perc = q.riga["conto_nome"], " > ".join(q.riga["percorso"])
        if "Ristorante Mare Mosso" in perc or "maremosso" in nome.lower():
            assert q.segmento_id == MMS


def test_segmento_riga_non_coperta(risultato):
    righe = _righe_db(risultato)
    senza_catch_all = REGOLE_PROVA[1:]
    quote, errori = assegna_segmenti(righe, senza_catch_all)
    assert errori
    assert all("nessuna regola" in e.motivo for e in errori)
    # Le righe in errore non sono in nessun segmento, ma sono riportate: nulla sparisce in silenzio
    assert len({id(q.riga) for q in quote}) + len(errori) == len(righe)


def test_segmento_percentuali_non_100(risultato):
    righe = _righe_db(risultato)
    regole = REGOLE_PROVA[:4]   # utenze: solo 70% assegnato
    _, errori = assegna_segmenti(righe, regole)
    assert errori and all("70" in e.motivo for e in errori)
    assert {e.riga["conto_nome"] for e in errori} == {"Acqua e Fognatura", "Elettricità", "Gas"}


def test_regola_validita_temporale(risultato):
    righe = _righe_db(risultato)
    regole = [
        _regola("percorso_prefisso", "", HOTEL, prio=0),
        {**_regola("percorso_prefisso", "Food & Beverage > Ristorante", MMS), "valida_dal": "2026-07"},
    ]
    quote, errori = assegna_segmenti(righe, regole)
    assert errori == []
    mms = _per_mese(quote, MMS)
    assert all(m >= 7 for (_, m) in mms)


# ── Export con subtotali sulle intestazioni (formato CLB) ────────────────────

def _con_subtotali(xlsx_bytes, errore_su=None):
    """Scrive su ogni intestazione la somma delle foglie sotto di lei, come fanno alcuni export
    del gestionale. errore_su: nome di un'intestazione su cui sbagliare il subtotale di 1 €."""
    def modifica(ws):
        righe = [(r, int(ws.cell(r, 1).alignment.indent or 0)) for r in range(4, ws.max_row)
                 if ws.cell(r, 1).value]   # esclusa l'ultima riga (Gran Totale)
        foglie = {r for i, (r, ind) in enumerate(righe)
                  if i + 1 == len(righe) or righe[i + 1][1] <= ind}
        for i, (r, ind) in enumerate(righe):
            if r in foglie:
                continue
            discendenti = []
            for r2, ind2 in righe[i + 1:]:
                if ind2 <= ind:
                    break
                if r2 in foglie:
                    discendenti.append(r2)
            for col in range(2, 11):
                somma = sum((ws.cell(f, col).value or 0) for f in discendenti)
                ws.cell(r, col).value = somma + (1 if ws.cell(r, 1).value == errore_su else 0)
    return _xlsx_modificato(xlsx_bytes, modifica)


def test_subtotali_su_intestazioni_accettati(xlsx_bytes):
    r = parse_export(_con_subtotali(xlsx_bytes))
    assert len(r.righe) == 223
    assert _totali(r.righe) == _d("-6831.55", "-7933.76", "-22907.78", "-44562.45", "-24407.67",
                                  "7971.45", "67716.00", "129358.28", "9556.32")


def test_subtotale_sbagliato_rifiutato(xlsx_bytes):
    with pytest.raises(ErroreImportUsali, match="Utenze.*non coincide"):
        parse_export(_con_subtotali(xlsx_bytes, errore_su="Utenze"))
