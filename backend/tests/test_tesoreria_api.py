"""Test integrazione modulo Tesoreria (blocco B): import estratti conto con anteprima/conferma,
antiduplicato su file sovrapposti, saldi, eliminazione import, movimenti (filtri, correzione
categoria, proposta di regola), configurazione (conti, saldi manuali, categorie, regole, prova,
ricalcolo) e permessi per modulo.

Fixture finte in tests/fixtures/tesoreria/. Seed di categorie e regole letto dalla migrazione
tes001_2026 (stesse costanti usate in produzione). DB dedicato, file in una cartella temporanea.
"""
import importlib.util
import json
import os
import re
import sys
from types import SimpleNamespace

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from sqlalchemy import create_engine, text  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402
from starlette.testclient import TestClient  # noqa: E402

from app.auth import richiedi_admin, richiedi_utente_attivo  # noqa: E402
from app.database import Base, get_db  # noqa: E402
from app.main import app  # noqa: E402

from tests._db import TEST_DB_URL  # noqa: E402

BACKEND = os.path.join(os.path.dirname(__file__), "..")
FIXTURES = os.path.join(os.path.dirname(__file__), "fixtures", "tesoreria")
IBAN_PRINCIPALE = "IT00X0000000000000000111111"
IBAN_POS = "IT00X0000000000000000222222"


def _leggi(nome):
    with open(os.path.join(FIXTURES, nome), "rb") as f:
        return f.read()


PRINCIPALE = _leggi("estratto_principale.xml")
POS = _leggi("estratto_pos.xml")
# Stesso conto senza il giorno 05/03 (primo <Stmt> del file): periodo sovrapposto, file diverso
PRINCIPALE_SENZA_05 = re.sub(rb"<ns5:Stmt>.*?</ns5:Stmt>", b"", PRINCIPALE, count=1, flags=re.S)
# Stesso estratto un anno prima (riferimenti banca diversi): confronto anno precedente
PRINCIPALE_2025 = PRINCIPALE.replace(b"2026-03", b"2025-03").replace(b"AcctSvcrRef>9", b"AcctSvcrRef>7")


def _migrazione():
    percorso = os.path.join(BACKEND, "alembic", "versions", "tes001_2026_modulo_tesoreria.py")
    spec = importlib.util.spec_from_file_location("tes001_2026", percorso)
    modulo = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modulo)
    return modulo


MIGRAZIONE = _migrazione()
ADMIN = SimpleNamespace(id=None, ruolo="admin")


@pytest.fixture(scope="module")
def test_engine():
    engine = create_engine(TEST_DB_URL)
    Base.metadata.create_all(engine)
    yield engine
    Base.metadata.drop_all(engine)
    engine.dispose()


@pytest.fixture(scope="module")
def client(test_engine, tmp_path_factory):
    os.environ["TESORERIA_DIR"] = str(tmp_path_factory.mktemp("tesoreria"))
    Sessione = sessionmaker(bind=test_engine)

    def override_get_db():
        db = Sessione()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[richiedi_utente_attivo] = lambda: utente_corrente["u"]
    app.dependency_overrides[richiedi_admin] = lambda: utente_corrente["u"]
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()
    os.environ.pop("TESORERIA_DIR", None)


utente_corrente = {"u": ADMIN}


@pytest.fixture(autouse=True)
def pulisci_e_semina(test_engine):
    utente_corrente["u"] = ADMIN
    with test_engine.begin() as conn:
        conn.execute(text(
            "TRUNCATE tesoreria_movimenti, tesoreria_saldi, tesoreria_import, tesoreria_regole, "
            "tesoreria_categorie, tesoreria_conti, user_module_permissions, module_permissions, modules, "
            "users, cassa_movimenti, app_config RESTART IDENTITY CASCADE"
        ))
        conn.execute(text(
            "INSERT INTO modules (code, name, ordine, attivo) VALUES ('tesoreria', 'TESORERIA', 30, true)"
        ))
        conn.execute(text("""
            INSERT INTO module_permissions (module_code, ruolo, puo_vedere, puo_modificare, puo_importare)
            VALUES ('tesoreria', 'admin', true, true, true), ('tesoreria', 'viewer', false, false, false)
        """))
        conn.execute(text("""
            INSERT INTO tesoreria_conti (codice, nome, tipo, iban, fido, fido_scadenza, ha_import, ordine,
                                         testo_addebito) VALUES
            ('PRINC', 'Conto principale', 'conto', :p, 0, NULL, true, 1, NULL),
            ('POS', 'Conto POS', 'conto', :pos, 130000, '2027-09-30', true, 2, NULL),
            ('AMEX', 'American Express', 'carta_credito', NULL, 0, NULL, false, 3, 'AMERICAN EXPRESS')
        """), {"p": IBAN_PRINCIPALE, "pos": IBAN_POS})
        for ordine, (codice, nome, tipo, escludi, colore) in enumerate(MIGRAZIONE.CATEGORIE, start=1):
            conn.execute(text(
                "INSERT INTO tesoreria_categorie (codice, nome, tipo, escludi_da_operativo, colore, ordine) "
                "VALUES (:c, :n, :t, :e, :col, :o)"
            ), dict(c=codice, n=nome, t=tipo, e=escludi, col=colore, o=ordine))
        for i, (cat, campo, op, valori, anche, segno) in enumerate(MIGRAZIONE.REGOLE, start=1):
            conn.execute(text("""
                INSERT INTO tesoreria_regole (categoria_id, priorita, campo, operatore, valori, valori_anche, segno)
                SELECT id, :pr, :campo, :op, CAST(:v AS JSON), CAST(:a AS JSON), :s
                FROM tesoreria_categorie WHERE codice = :cat
            """), dict(cat=cat, pr=i * 10, campo=campo, op=op, v=json.dumps(valori),
                       a=json.dumps(anche) if anche else None, s=segno))


def _anteprima(client, contenuto=PRINCIPALE, nome="estratto.xml"):
    return client.post("/tesoreria/import/anteprima", files={"file": (nome, contenuto, "application/xml")})


def _importa(client, contenuto=PRINCIPALE, nome="estratto.xml"):
    r = _anteprima(client, contenuto, nome)
    assert r.status_code == 200, r.text
    c = client.post("/tesoreria/import/conferma", json={"token": r.json()["token"]})
    assert c.status_code == 200, c.text
    return c.json()


def _movimenti(client, **params):
    r = client.get("/tesoreria/movimenti", params=params)
    assert r.status_code == 200, r.text
    return r.json()


def _id_categoria(client, codice):
    return next(c["id"] for c in client.get("/tesoreria/categorie").json() if c["codice"] == codice)


# ── Import ──────────────────────────────────────────────────────────────────────

def test_anteprima_riconosce_conto_e_non_scrive(client):
    r = _anteprima(client)
    assert r.status_code == 200, r.text
    a = r.json()
    assert a["conto"]["codice"] == "PRINC"
    assert (a["data_da"], a["data_a"]) == ("2026-03-02", "2026-03-05")
    assert (a["n_movimenti"], a["n_nuovi"], a["n_gia_presenti"]) == (7, 7, 0)
    assert a["saldo_finale"] == 9820.61
    assert a["discrepanze"] == [] and a["gia_importato"] is None
    per_cat = {c["codice"]: c["n"] for c in a["per_categoria"]}
    assert per_cat == {"bonifici_clienti": 1, "banca": 2, "fornitori_bonifici": 1, "giroconti": 1,
                       "mutuo_rate": 1, "assegni": 1}
    assert _movimenti(client)["totale"] == 0


def test_anteprima_iban_sconosciuto(client):
    xml = PRINCIPALE.replace(IBAN_PRINCIPALE.encode(), b"IT00X0000000000000000999999")
    r = _anteprima(client, xml)
    assert r.status_code == 400 and "IT00X0000000000000000999999" in r.json()["detail"]


def test_anteprima_file_non_valido(client):
    r = _anteprima(client, b"non sono un xml")
    assert r.status_code == 400 and "XML" in r.json()["detail"]


def test_conferma_inserisce_movimenti_saldi_e_file(client):
    esito = _importa(client)
    assert (esito["n_nuovi"], esito["n_saldi"], esito["n_discrepanze"]) == (7, 4, 0)
    dati = _movimenti(client)
    assert dati["totale"] == 7 and dati["netto"] == -179.39
    saldi = client.get("/tesoreria/saldi", params={"conto_id": 1}).json()
    assert [(s["data"], s["saldo"], s["origine"]) for s in saldi][:2] == [
        ("2026-03-05", 9820.61, "import"), ("2026-03-03", 10820.61, "import")]
    storico = client.get("/tesoreria/import").json()
    assert storico[0]["n_collegati"] == 7 and storico[0]["ha_file"]
    f = client.get(f"/tesoreria/import/{esito['import_id']}/file")
    assert f.status_code == 200 and f.content == PRINCIPALE
    conto = next(c for c in client.get("/tesoreria/conti").json() if c["codice"] == "PRINC")
    assert conto["ultimo_saldo"] == {"data": "2026-03-05", "saldo": 9820.61, "origine": "import"}


def test_token_usato_una_volta(client):
    token = _anteprima(client).json()["token"]
    assert client.post("/tesoreria/import/conferma", json={"token": token}).status_code == 200
    assert client.post("/tesoreria/import/conferma", json={"token": token}).status_code == 410


def test_stesso_file_rifiutato(client):
    _importa(client)
    a = _anteprima(client).json()
    assert a["gia_importato"] is not None and a["n_nuovi"] == 0
    c = client.post("/tesoreria/import/conferma", json={"token": a["token"]})
    assert c.status_code == 409


def test_file_sovrapposto_nessun_doppione(client):
    _importa(client, PRINCIPALE_SENZA_05, "parziale.xml")
    assert _movimenti(client)["totale"] == 6
    a = _anteprima(client).json()
    assert (a["n_nuovi"], a["n_gia_presenti"]) == (1, 6)
    _importa(client)
    assert _movimenti(client)["totale"] == 7
    assert client.get("/tesoreria/import").json()[0]["n_collegati"] == 1


def test_discrepanza_saldo_segnalata(client, test_engine):
    """Un movimento mancante tra due saldi noti (= file o periodo mancante) fa scattare la verifica."""
    _importa(client)
    with test_engine.begin() as conn:
        # l'assegno del 05/03, che il file parziale (fino al 03/03) non contiene
        conn.execute(text("DELETE FROM tesoreria_movimenti WHERE descrizione = 'Vs assegno n.'"))
    a = _anteprima(client, PRINCIPALE_SENZA_05, "parziale.xml").json()
    assert a["n_nuovi"] == 0
    assert len(a["discrepanze"]) == 1
    d = a["discrepanze"][0]
    assert (d["data_da"], d["data_a"], d["differenza"]) == ("2026-03-03", "2026-03-05", -1000.0)


def test_estratto_sostituisce_saldo_manuale(client):
    r = client.post("/tesoreria/saldi", json={"conto_id": 1, "data": "2026-03-03", "saldo": 10000})
    assert r.status_code == 200
    a = _anteprima(client).json()
    assert a["saldi_cambiati"] == [{"data": "2026-03-03", "saldo_attuale": 10000.0, "origine": "manuale",
                                    "saldo_nuovo": 10820.61}]
    _importa(client)
    saldo = next(s for s in client.get("/tesoreria/saldi").json() if s["data"] == "2026-03-03")
    assert (saldo["saldo"], saldo["origine"]) == (10820.61, "import")


def test_elimina_import(client):
    primo = _importa(client, PRINCIPALE_SENZA_05, "parziale.xml")
    secondo = _importa(client)
    assert client.delete(f"/tesoreria/import/{secondo['import_id']}").status_code == 400
    r = client.delete(f"/tesoreria/import/{secondo['import_id']}", params={"conferma": True})
    assert r.status_code == 200
    # Toglie solo ciò che quell'import aveva inserito per primo: assegno e saldo del 05/03 (gli altri
    # saldi del file erano già presenti e uguali, restano del primo import)
    assert (r.json()["movimenti_eliminati"], r.json()["saldi_eliminati"]) == (1, 1)
    assert _movimenti(client)["totale"] == 6
    assert [i["id"] for i in client.get("/tesoreria/import").json()] == [primo["import_id"]]


# ── Movimenti ───────────────────────────────────────────────────────────────────

def test_filtri_ricerca_e_ordinamento(client):
    _importa(client)
    _importa(client, POS, "pos.xml")
    assert _movimenti(client)["totale"] == 11
    assert _movimenti(client, conto_id=2)["totale"] == 4
    assert _movimenti(client, segno="entrata", conto_id=1)["totale"] == 2
    assert _movimenti(client, data_da="2026-03-03", data_a="2026-03-03", conto_id=1)["totale"] == 2
    assert _movimenti(client, q="fornitore alfa")["totale"] == 1
    assert _movimenti(client, q="1478")["totale"] == 1
    assert _movimenti(client, categoria_id=_id_categoria(client, "giroconti"))["totale"] == 2
    righe = _movimenti(client, ordina_per="importo", direzione="asc")["righe"]
    assert righe[0]["importo"] == -3000.0
    pagina = _movimenti(client, per_pagina=5, pagina=3)
    assert len(pagina["righe"]) == 1 and pagina["totale"] == 11


def test_correzione_manuale_resiste_al_ricalcolo(client):
    _importa(client)
    mov = _movimenti(client, q="1478")["righe"][0]
    assert (mov["categoria_codice"], mov["categoria_manuale"]) == ("mutuo_rate", False)
    banca = _id_categoria(client, "banca")
    r = client.put(f"/tesoreria/movimenti/{mov['id']}",
                   json={"categoria_id": banca, "straordinario": True, "note": " verificato "})
    assert r.status_code == 200
    m = r.json()
    assert (m["categoria_codice"], m["categoria_manuale"], m["straordinario"], m["note"]) == \
        ("banca", True, True, "verificato")
    assert client.post("/tesoreria/ricalcola").json()["n_cambiati"] == 0
    assert _movimenti(client, q="1478")["righe"][0]["categoria_codice"] == "banca"
    r = client.put(f"/tesoreria/movimenti/{mov['id']}", json={"automatica": True})
    assert (r.json()["categoria_codice"], r.json()["categoria_manuale"]) == ("mutuo_rate", False)
    assert client.put(f"/tesoreria/movimenti/{mov['id']}", json={"categoria_id": 9999}).status_code == 400


def test_proposta_regola_e_creazione_con_ricalcolo(client):
    _importa(client)
    mov = _movimenti(client, q="fornitore alfa")["righe"][0]
    assert mov["categoria_codice"] == "fornitori_bonifici"
    p = client.get(f"/tesoreria/movimenti/{mov['id']}/proposta-regola").json()
    assert p["campo"] == "controparte" and p["operatore"] == "inizia" and p["segno"] == "uscita"
    assert p["valori"][0].startswith("FORNITORE ALFA SRL")
    regola_attuale = next(r for r in client.get("/tesoreria/regole").json() if r["id"] == p["regola_attuale_id"])
    assert p["priorita"] == regola_attuale["priorita"] - 1

    utenze = _id_categoria(client, "utenze")
    bozza = {**p, "categoria_id": utenze, "valori": ["FORNITORE ALFA"]}
    prova = client.post("/tesoreria/regole/prova", json=bozza).json()
    assert (prova["n_corrispondenti"], prova["n_cambierebbero"]) == (1, 1)
    assert prova["da_categorie"] == [{"categoria": "Fornitori – bonifici", "n": 1}]
    assert _movimenti(client, q="fornitore alfa")["righe"][0]["categoria_codice"] == "fornitori_bonifici"

    r = client.post("/tesoreria/regole", json={**bozza, "ricalcola": True})
    assert r.status_code == 200 and r.json()["ricalcolo"]["n_cambiati"] == 1
    assert _movimenti(client, q="fornitore alfa")["righe"][0]["categoria_codice"] == "utenze"

    r = client.delete(f"/tesoreria/regole/{r.json()['id']}", params={"ricalcola": True})
    assert r.json()["ricalcolo"]["n_cambiati"] == 1
    assert _movimenti(client, q="fornitore alfa")["righe"][0]["categoria_codice"] == "fornitori_bonifici"


def test_regola_non_valida(client):
    utenze = _id_categoria(client, "utenze")
    assert client.post("/tesoreria/regole", json={"categoria_id": utenze, "valori": []}).status_code == 400
    assert client.post("/tesoreria/regole", json={"categoria_id": utenze, "valori": ["X"],
                                                  "operatore": "uguale"}).status_code == 400


def test_movimento_non_classificato_va_in_ripiego(client, test_engine):
    """Senza regole che corrispondono → 'Da classificare' (mai NULL)."""
    with test_engine.begin() as conn:
        conn.execute(text("DELETE FROM tesoreria_regole"))
    _importa(client)
    ripiego = _id_categoria(client, "da_classificare")
    assert _movimenti(client, categoria_id=ripiego)["totale"] == 7


# ── Configurazione ──────────────────────────────────────────────────────────────

def test_conti_fido_e_validazioni(client):
    r = client.put("/tesoreria/conti/2", json={"fido": "150000", "fido_scadenza": "2028-09-30"})
    assert r.status_code == 200
    assert (r.json()["fido"], r.json()["fido_scadenza"]) == (150000.0, "2028-09-30")
    assert client.put("/tesoreria/conti/2", json={"fido": -1}).status_code == 400
    assert client.put("/tesoreria/conti/2", json={"iban": IBAN_PRINCIPALE}).status_code == 409
    assert client.post("/tesoreria/conti", json={"codice": "X", "nome": "X", "tipo": "conto",
                                                 "ha_import": True}).status_code == 400
    r = client.post("/tesoreria/conti", json={"codice": "CASSA", "nome": "Cassa", "tipo": "cassa"})
    assert r.status_code == 200 and r.json()["ha_import"] is False


def test_saldi_manuali_carte(client):
    r = client.post("/tesoreria/saldi", json={"conto_id": 3, "data": "2026-03-06", "saldo": "3042.39"})
    assert r.status_code == 200
    r = client.post("/tesoreria/saldi", json={"conto_id": 3, "data": "2026-03-06", "saldo": 3100, "note": "x"})
    assert r.json()["saldo"] == 3100.0
    assert len(client.get("/tesoreria/saldi", params={"conto_id": 3}).json()) == 1
    assert client.delete(f"/tesoreria/saldi/{r.json()['id']}").status_code == 200


def test_saldo_da_estratto_non_modificabile(client):
    _importa(client)
    r = client.post("/tesoreria/saldi", json={"conto_id": 1, "data": "2026-03-05", "saldo": 1})
    assert r.status_code == 409
    saldo = client.get("/tesoreria/saldi", params={"conto_id": 1}).json()[0]
    assert client.delete(f"/tesoreria/saldi/{saldo['id']}").status_code == 409


def test_categorie(client):
    _importa(client)
    r = client.post("/tesoreria/categorie", json={"codice": "Nuova Cat", "nome": "Nuova", "tipo": "uscita",
                                                  "colore": "#AABBCC"})
    assert r.status_code == 200 and (r.json()["codice"], r.json()["colore"]) == ("nuova_cat", "#aabbcc")
    assert client.delete(f"/tesoreria/categorie/{r.json()['id']}").status_code == 200
    banca = _id_categoria(client, "banca")
    assert client.delete(f"/tesoreria/categorie/{banca}").status_code == 409
    ripiego = _id_categoria(client, "da_classificare")
    assert client.put(f"/tesoreria/categorie/{ripiego}", json={"attivo": False}).status_code == 400
    assert client.delete(f"/tesoreria/categorie/{ripiego}").status_code == 400
    assert client.put(f"/tesoreria/categorie/{banca}", json={"colore": "rosso"}).status_code == 400
    elenco = {c["codice"]: c for c in client.get("/tesoreria/categorie").json()}
    assert elenco["banca"]["n_movimenti"] == 2 and elenco["banca"]["n_regole"] > 0
    assert elenco["da_classificare"]["protetta"]


def test_categoria_disattivata_esclusa_dalle_regole(client):
    _importa(client)
    mutuo = _id_categoria(client, "mutuo_rate")
    client.put(f"/tesoreria/categorie/{mutuo}", json={"attivo": False})
    client.post("/tesoreria/ricalcola")
    assert _movimenti(client, q="1478")["righe"][0]["categoria_codice"] != "mutuo_rate"


# ── Permessi ────────────────────────────────────────────────────────────────────

def _viewer(test_engine, **permessi):
    with test_engine.begin() as conn:
        conn.execute(text("INSERT INTO users (username, password_hash, ruolo) VALUES ('ospite', 'x', 'viewer')"))
        if permessi:
            sets = ", ".join(f"{k} = :{k}" for k in permessi)
            conn.execute(text(f"UPDATE module_permissions SET {sets} "
                              "WHERE module_code = 'tesoreria' AND ruolo = 'viewer'"), permessi)
    utente_corrente["u"] = SimpleNamespace(id=1, ruolo="viewer")


def test_viewer_senza_permessi_403(client, test_engine):
    _viewer(test_engine)
    assert client.get("/tesoreria/movimenti").status_code == 403
    assert client.get("/tesoreria/conti").status_code == 403
    assert _anteprima(client).status_code == 403


def test_viewer_abilitato_solo_lettura(client, test_engine):
    _viewer(test_engine, puo_vedere=True)
    assert client.get("/tesoreria/movimenti").status_code == 200
    assert client.post("/tesoreria/ricalcola").status_code == 403
    assert _anteprima(client).status_code == 403


def test_viewer_override_per_utente(client, test_engine):
    """L'override per utente vale solo per 'vedere', come nella NavBar."""
    _viewer(test_engine, puo_modificare=True)
    assert client.post("/tesoreria/ricalcola").status_code == 403    # modifica senza vedere: no
    with test_engine.begin() as conn:
        conn.execute(text("INSERT INTO user_module_permissions (user_id, module_code, puo_vedere) "
                          "VALUES (1, 'tesoreria', true)"))
    assert client.get("/tesoreria/movimenti").status_code == 200
    assert client.post("/tesoreria/ricalcola").status_code == 200
    moduli = client.get("/modules/").json()
    assert next(m for m in moduli if m["code"] == "tesoreria")["puo_vedere"] is True


def test_modulo_disattivato_solo_admin(client, test_engine):
    _viewer(test_engine, puo_vedere=True)
    with test_engine.begin() as conn:
        conn.execute(text("UPDATE modules SET attivo = false WHERE code = 'tesoreria'"))
    assert client.get("/tesoreria/movimenti").status_code == 403
    utente_corrente["u"] = ADMIN
    assert client.get("/tesoreria/movimenti").status_code == 200


# ── Posizione attuale ───────────────────────────────────────────────────────────

def _posizione(client, giorno):
    r = client.get("/tesoreria/posizione", params={"giorno": giorno})
    assert r.status_code == 200, r.text
    return r.json()


def _cassa(test_engine, importo, data="2026-01-01"):
    with test_engine.begin() as conn:
        conn.execute(text("INSERT INTO cassa_movimenti (tipo, data, importo, is_test) "
                          "VALUES ('saldo_iniziale', :d, :i, false)"), {"d": data, "i": importo})


def test_posizione_voci_e_disponibilita(client, test_engine):
    _importa(client)
    _importa(client, POS, "pos.xml")
    _cassa(test_engine, 1000)
    client.post("/tesoreria/saldi", json={"conto_id": 3, "data": "2026-03-01", "saldo": 500})
    p = _posizione(client, "2026-03-06")
    assert p["saldo_banche"] == 9682.41                       # 9.820,61 + (−138,20)
    assert (p["contanti"], p["contanti_data_riferimento"]) == (1000.0, "2026-01-01")
    assert p["carte_da_addebitare"] == -500.0
    assert p["liquidita_netta"] == 10182.41
    assert (p["fidi"], p["disponibilita_con_fido"]) == (130000.0, 140182.41)
    pos = next(c for c in p["conti"] if c["codice"] == "POS")
    assert (pos["saldo"], pos["data_ultimo_saldo"], pos["fido_conteggiato"]) == (-138.2, "2026-03-04", True)
    assert [a for a in p["avvisi"] if a["livello"] != "info"] == []


def test_posizione_saldo_banca_a_data_intermedia(client):
    _importa(client)
    assert _posizione(client, "2026-03-04")["saldo_banche"] == 10820.61   # nessun movimento il 04/03


def test_posizione_fido_in_scadenza_e_scaduto(client):
    _importa(client, POS, "pos.xml")
    p = _posizione(client, "2027-08-15")
    assert p["fidi"] == 130000.0
    assert any(a["tipo"] == "fido" and a["livello"] == "arancione" for a in p["avvisi"])
    p = _posizione(client, "2027-10-01")
    assert p["fidi"] == 0.0 and p["disponibilita_con_fido"] == p["liquidita_netta"]
    assert any(a["tipo"] == "fido" and a["livello"] == "rosso" for a in p["avvisi"])
    assert p["avvisi"][0]["livello"] == "rosso"


def test_posizione_avvisi_freschezza(client, test_engine):
    client.post("/tesoreria/saldi", json={"conto_id": 3, "data": "2026-01-01", "saldo": 100})
    p = _posizione(client, "2026-03-06")
    tipi = {(a["tipo"], a["conto_id"]) for a in p["avvisi"]}
    assert ("estratto", 1) in tipi and ("estratto", 2) in tipi           # nessun estratto caricato
    assert ("saldo_manuale", 3) in tipi                                   # 64 giorni > 35
    assert ("contanti", None) in tipi                                     # Cassa senza dati
    client.put("/tesoreria/soglie", json={"giorni_saldo_carta": 90})
    assert ("saldo_manuale", 3) not in {(a["tipo"], a["conto_id"]) for a in _posizione(client, "2026-03-06")["avvisi"]}


def test_posizione_segnala_discrepanza(client, test_engine):
    _importa(client)
    with test_engine.begin() as conn:
        conn.execute(text("DELETE FROM tesoreria_movimenti WHERE descrizione = 'Vs assegno n.'"))
    rossi = [a for a in _posizione(client, "2026-03-06")["avvisi"] if a["tipo"] == "discrepanza"]
    assert len(rossi) == 1 and "-1.000,00 €" in rossi[0]["messaggio"]


def test_soglie(client):
    assert client.get("/tesoreria/soglie").json()["giorni_estratto"] == 7
    r = client.put("/tesoreria/soglie", json={"giorni_estratto": 10})
    assert r.json() == {"giorni_estratto": 10, "giorni_saldo_carta": 35, "giorni_preavviso_fido": 60}
    assert client.put("/tesoreria/soglie", json={"giorni_estratto": 0}).status_code == 400
    assert client.get("/tesoreria/soglie").json()["giorni_estratto"] == 10


# ── Consuntivo ──────────────────────────────────────────────────────────────────

def _consuntivo(client, **params):
    r = client.get("/tesoreria/consuntivo", params={"data_da": "2026-03-01", "data_a": "2026-03-31", **params})
    assert r.status_code == 200, r.text
    return r.json()


def test_consuntivo_tutti_i_conti_senza_giroconti(client):
    _importa(client)
    _importa(client, POS, "pos.xml")
    c = _consuntivo(client)
    k = c["kpi"]
    assert (k["entrate"]["valore"], k["uscite"]["valore"], k["netto_operativo"]["valore"]) == \
        (3700.0, -4517.59, -817.59)
    assert (k["saldo_inizio"]["valore"], k["saldo_fine"]["valore"]) == (10500.0, 9682.41)
    assert c["giroconti_esclusi"] == {"n": 2, "entrate": 3000.0, "uscite": -3000.0, "netto": 0.0}
    assert c["quadratura"]["ok"]
    assert "giroconti" not in {r["codice"] for r in c["categorie"]}
    assert c["mensile"] == [{"anno": 2026, "mese": 3, "etichetta": "mar 2026", "entrate": 3700.0,
                             "uscite": -4517.59, "netto_operativo": -817.59, "netto_operativo_ap": 0.0}]
    assert c["ap_completo"] is False and k["entrate"]["anno_precedente"] == 0.0


def test_consuntivo_singolo_conto_giroconti_non_operativi(client):
    _importa(client)
    c = _consuntivo(client, conto_id=1)
    righe = {r["codice"]: r for r in c["categorie"]}
    assert righe["giroconti"]["operativa"] is False and righe["giroconti"]["netto"] == 3000.0
    assert c["kpi"]["non_operativi"]["valore"] == 3000.0
    assert c["kpi"]["netto_operativo"]["valore"] == -3179.39
    assert c["quadratura"]["ok"] and c["giroconti_esclusi"]["n"] == 0
    # operative prima, non operative in fondo
    assert [r["operativa"] for r in c["categorie"]] == sorted((r["operativa"] for r in c["categorie"]), reverse=True)


def test_consuntivo_anno_precedente(client):
    _importa(client, PRINCIPALE_2025, "2025.xml")
    _importa(client)
    c = _consuntivo(client, conto_id=1)
    assert c["ap_completo"] is True
    k = c["kpi"]
    assert k["netto_operativo"]["anno_precedente"] == -3179.39 and k["netto_operativo"]["delta"] == 0.0
    assert k["saldo_fine"]["anno_precedente"] == 9820.61
    assert c["mensile"][0]["netto_operativo_ap"] == -3179.39
    banca = next(r for r in c["categorie"] if r["codice"] == "banca")
    assert (banca["netto"], banca["netto_ap"], banca["delta_pct"]) == (-0.5, -0.5, 0.0)


def test_consuntivo_periodo_e_conto_non_validi(client):
    r = client.get("/tesoreria/consuntivo", params={"data_da": "2026-03-31", "data_a": "2026-03-01"})
    assert r.status_code == 400
    assert client.get("/tesoreria/consuntivo", params={"data_da": "2026-03-01", "data_a": "2026-03-31",
                                                       "conto_id": 3}).status_code == 404


def test_posizione_carta_torna_a_zero_dopo_addebito(client, test_engine):
    """Si inserisce solo il totale dell'estratto alla chiusura: la carta torna a zero quando la banca
    addebita lo stesso importo."""
    _importa(client)
    client.post("/tesoreria/saldi", json={"conto_id": 3, "data": "2026-02-28", "saldo": "1200.00"})
    with test_engine.begin() as conn:   # addebito AMEX finto sul conto principale il 04/03
        conn.execute(text("""
            INSERT INTO tesoreria_movimenti (conto_id, data_contabile, importo, descrizione, impronta, categoria_manuale,
                                             straordinario)
            VALUES (1, '2026-03-04', -1200.00, 'SDD Core - Richiesta Incasso SEPA 01M4 AMERICAN EXPRESS ITALIA',
                    'R:test-amex', false, false)
        """))
    prima = _posizione(client, "2026-03-03")
    amex = next(c for c in prima["conti"] if c["codice"] == "AMEX")
    assert prima["carte_da_addebitare"] == -1200.0
    assert (amex["carta_stato"], amex["carta_importo_estratto"]) == ("da_addebitare", 1200.0)
    dopo = _posizione(client, "2026-03-06")
    amex = next(c for c in dopo["conti"] if c["codice"] == "AMEX")
    assert dopo["carte_da_addebitare"] == 0.0
    assert (amex["carta_stato"], amex["carta_data_addebito"], amex["saldo"]) == ("addebitato", "2026-03-04", 0.0)
    # 40 giorni dopo la chiusura: promemoria per l'estratto successivo
    avviso = next(a for a in _posizione(client, "2026-04-09")["avvisi"] if a["conto_id"] == 3)
    assert "estratto successivo" in avviso["messaggio"]


def test_posizione_carta_addebito_mancante(client):
    client.post("/tesoreria/saldi", json={"conto_id": 3, "data": "2026-02-28", "saldo": "1200.00"})
    avviso = next(a for a in _posizione(client, "2026-04-09")["avvisi"] if a["conto_id"] == 3)
    assert "non risulta ancora addebitato" in avviso["messaggio"] and "1.200,00 €" in avviso["messaggio"]


# ── Export ──────────────────────────────────────────────────────────────────────

def test_export_movimenti_con_filtri(client):
    import io
    import openpyxl
    _importa(client)
    r = client.get("/tesoreria/movimenti/export", params={"formato": "xlsx", "segno": "uscita"})
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    righe = list(ws.iter_rows(values_only=True))
    assert righe[0][:6] == ("Data contabile", "Data valuta", "Conto", "Controparte", "Descrizione", "Importo")
    assert len(righe) == 1 + 5                      # 5 uscite nella fixture
    assert righe[1][0] == "05/03/2026" and righe[1][5] == -1000.0   # ordine per data decrescente
    csv = client.get("/tesoreria/movimenti/export", params={"formato": "csv", "q": "1478"}).content.decode("utf-8-sig")
    assert "-1478,89" in csv and csv.count("\n") == 2
    pdf = client.get("/tesoreria/movimenti/export", params={"formato": "pdf"})
    assert pdf.content[:4] == b"%PDF"


def test_export_consuntivo(client):
    import io
    import openpyxl
    _importa(client)
    _importa(client, POS, "pos.xml")
    r = client.get("/tesoreria/consuntivo/export",
                   params={"data_da": "2026-03-01", "data_a": "2026-03-31", "formato": "xlsx"})
    assert r.status_code == 200
    righe = list(openpyxl.load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
    per_voce = {riga[0]: riga for riga in righe}
    assert per_voce["TOTALE OPERATIVO"][2:5] == (3700.0, -4517.59, -817.59)
    assert per_voce["Saldo fine periodo"][4] == 9682.41
    assert per_voce["mar 2026"][4] == -817.59
    pdf = client.get("/tesoreria/consuntivo/export",
                     params={"data_da": "2026-03-01", "data_a": "2026-03-31", "formato": "pdf"})
    assert pdf.content[:4] == b"%PDF"


def test_riprocessa_dai_file_originali(client, test_engine):
    """Controparte e causale si ricalcolano dai file XML conservati; categorie manuali invariate."""
    _importa(client)
    mutuo = _movimenti(client, q="1478")["righe"][0]
    client.put(f"/tesoreria/movimenti/{mutuo['id']}", json={"categoria_id": _id_categoria(client, "banca")})
    with test_engine.begin() as conn:
        conn.execute(text("UPDATE tesoreria_movimenti SET controparte = 'VECCHIO', causale = NULL"))
    r = client.post("/tesoreria/import/riprocessa")
    assert r.status_code == 200, r.text
    esito = r.json()
    assert (esito["n_movimenti"], esito["n_riletti"], esito["n_aggiornati"], esito["n_senza_file"]) == (7, 7, 7, 0)
    righe = {m["importo"]: m for m in _movimenti(client)["righe"]}
    assert righe[500.0]["controparte"] == "MARIO ROSSI"
    assert righe[-1478.89]["controparte"] is None and righe[-1478.89]["causale"].startswith("Addebito rata mutuo")
    assert righe[-1478.89]["categoria_codice"] == "banca"          # manuale: non toccata


def test_kpi(client):
    _importa(client)
    _importa(client, POS, "pos.xml")
    r = client.get("/tesoreria/kpi", params={"anno": 2026})
    assert r.status_code == 200, r.text
    k = r.json()
    giorni = {p["data"]: p["saldo"] for p in k["andamento"]}
    assert giorni["2026-03-02"] == 9299.5 - 1300.0 and giorni["2026-03-04"] == 10820.61 - 138.2
    assert giorni["2026-01-01"] is None                       # prima dei dati: mancante, non zero
    assert k["minimo"]["anno"] == {"data": "2026-03-02", "saldo": 7999.5}
    pos = next(f for f in k["fido"] if f["nome"] == "Conto POS")
    assert pos["anno"]["saldo_minimo"] == -1300.0 and pos["anno"]["utilizzo_pct"] == 1.0
    assert pos["anno"]["giorni_sotto_zero"] >= 3
    nomi = [p["nome"] for p in k["peso_uscite"]]
    assert "Incassi POS/online" not in nomi and nomi[0] == "Mutuo – rate e spese"
    assert k["costo_incassi"]["anno"]["incassi_pos"] == 2361.8    # al netto dello storno
    assert k["tassa_soggiorno"]["disponibile"] is False
